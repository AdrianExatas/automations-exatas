#!/usr/bin/env python3
"""Upsert de entradas na Lista Documental Mestra (formato FORM.QUA.003)."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl é necessário para update_lista_mestra.py") from exc


SHEET_NAME = "Lista Mestra Doc. Institucional"
OUTPUT_NAME = "FORM.QUA.003 - Lista Documental Mestra Aplicada.xlsx"

HEADERS = [
    "Nº",
    "Código e Título do Documento",
    "Origem",
    "Tipo",
    "Setor",
    "Nº Rev.",
    "Última",
    "Prazo (Ano)",
    "Elaborado Por",
    "Verificado Por",
    "Aprovado Por",
    "Data Aprov.",
    "Localização em Diretório",
    "Procedimento Raiz",
    "Procedimentos Citados",
    "Observações",
]

COL = {name: idx + 1 for idx, name in enumerate(HEADERS)}


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def code_key(codigo_titulo: str) -> str:
    return re.split(r"\s*-\s*", codigo_titulo.strip(), maxsplit=1)[0].strip().upper()


def ensure_template(template_path: Path) -> None:
    if template_path.is_file():
        return
    template_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws["A1"] = "LISTA DOCUMENTAL MESTRA APLICADA"
    ws["A1"].font = Font(bold=True, size=14)
    ws["T1"] = "Código:"
    ws["U1"] = "FORM.QUA.003"
    ws["T2"] = "Emissão:"
    ws["T3"] = "Versão:"
    ws["U3"] = "00"
    ws["T4"] = "Revisão:"
    for idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=6, column=idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A7"
    wb.save(template_path)


def find_sheet(wb):
    if SHEET_NAME in wb.sheetnames:
        return wb[SHEET_NAME]
    return wb[wb.sheetnames[0]]


def detect_header_row(ws) -> int:
    for row in range(1, min(20, ws.max_row + 1)):
        values = [text(ws.cell(row=row, column=c).value).lower() for c in range(1, 8)]
        joined = " | ".join(values)
        if "código" in joined and "título" in joined and "tipo" in joined:
            return row
        if "codigo" in joined and "titulo" in joined and "tipo" in joined:
            return row
    return 6


def map_columns(ws, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, min(ws.max_column, 40) + 1):
        raw = text(ws.cell(row=header_row, column=col).value).lower()
        if not raw:
            continue
        if raw in {"nº", "n°", "no", "nº.", "n."} or raw.startswith("nº"):
            mapping.setdefault("Nº", col)
        elif "código" in raw or "codigo" in raw:
            mapping.setdefault("Código e Título do Documento", col)
        elif raw == "origem":
            mapping.setdefault("Origem", col)
        elif raw == "tipo":
            mapping.setdefault("Tipo", col)
        elif raw == "setor":
            mapping.setdefault("Setor", col)
        elif "elabor" in raw:
            mapping.setdefault("Elaborado Por", col)
        elif "verific" in raw:
            mapping.setdefault("Verificado Por", col)
        elif "aprovado" in raw:
            mapping.setdefault("Aprovado Por", col)
        elif "data aprov" in raw:
            mapping.setdefault("Data Aprov.", col)
        elif "localiza" in raw:
            mapping.setdefault("Localização em Diretório", col)
        elif "procedimento raiz" in raw:
            mapping.setdefault("Procedimento Raiz", col)
        elif "procedimentos citados" in raw or "procedimentos cit" in raw:
            mapping.setdefault("Procedimentos Citados", col)
        elif "observa" in raw:
            mapping.setdefault("Observações", col)
        elif raw in {"nº rev.", "nº", "numero"} and "rev" in raw:
            mapping.setdefault("Nº Rev.", col)
        elif "última" in raw or "ultima" in raw:
            mapping.setdefault("Última", col)
    # Fallback to canonical positions for sanitized template.
    for name, idx in COL.items():
        mapping.setdefault(name, idx)
    return mapping


def iter_data_rows(ws, header_row: int, code_col: int) -> list[int]:
    rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        value = text(ws.cell(row=row, column=code_col).value)
        if value:
            rows.append(row)
    return rows


def next_empty_row(ws, header_row: int, code_col: int) -> int:
    row = header_row + 1
    while row <= ws.max_row and text(ws.cell(row=row, column=code_col).value):
        row += 1
    return row


def build_entries_from_content(content: dict[str, Any]) -> list[dict[str, Any]]:
    lista = content.get("lista_mestra") or {}
    entries = as_list(lista.get("entradas"))
    if entries:
        return [e for e in entries if isinstance(e, dict)]

    # Derive if normalization did not embed lista_mestra yet.
    doc = content.get("documento") or {}
    saida = content.get("saida") or {}
    types = {str(t).lower() for t in as_list(content.get("documentos_solicitados"))}
    title = text(doc.get("titulo")) or "PROCESSO"
    mapping = {
        "pop": ("PR", text(doc.get("codigo_pop")), text(saida.get("arquivo_pop"))),
        "it": ("IN", text(doc.get("codigo_it")), text(saida.get("arquivo_it"))),
        "form": ("FORM", text(doc.get("codigo_form")), text(saida.get("arquivo_form"))),
        "mp": ("MP", text(doc.get("codigo_mp")), text(saida.get("arquivo_mp"))),
    }
    raiz = text(doc.get("codigo_pop")) if "pop" in types else ""
    pack_codes = [mapping[t][1] for t in ("pop", "it", "form", "mp") if t in types and mapping[t][1]]
    derived = []
    for t in ("pop", "it", "form", "mp"):
        if t not in types:
            continue
        tipo, codigo, arquivo = mapping[t]
        citados = [c for c in pack_codes if c and c != codigo and c != raiz]
        derived.append(
            {
                "codigo_titulo": f"{codigo} - {title}",
                "origem": "INTERNO",
                "tipo": tipo,
                "setor": text(doc.get("setor")),
                "revisao_numero": text(doc.get("versao")) or "00",
                "ultima_revisao": text(doc.get("revisao")),
                "elaborador": text(doc.get("elaborador")),
                "verificador": text(doc.get("verificador")),
                "aprovador": text(doc.get("aprovador")),
                "data_aprovacao": text(doc.get("revisao")),
                "localizacao": f"documentos/{arquivo}" if arquivo else "",
                "procedimento_raiz": raiz,
                "procedimentos_citados": citados,
                "observacoes": "",
            }
        )
    return derived


def upsert(ws, entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    header_row = detect_header_row(ws)
    cols = map_columns(ws, header_row)
    code_col = cols["Código e Título do Documento"]
    existing_rows = {}
    for row in iter_data_rows(ws, header_row, code_col):
        key = code_key(text(ws.cell(row=row, column=code_col).value))
        if key:
            existing_rows[key] = row

    applied = []
    for entry in entries:
        codigo_titulo = text(entry.get("codigo_titulo"))
        if not codigo_titulo:
            continue
        key = code_key(codigo_titulo)
        row = existing_rows.get(key)
        created = False
        if row is None:
            row = next_empty_row(ws, header_row, code_col)
            existing_rows[key] = row
            created = True

        values = {
            "Código e Título do Documento": codigo_titulo,
            "Origem": text(entry.get("origem")) or "INTERNO",
            "Tipo": text(entry.get("tipo")),
            "Setor": text(entry.get("setor")),
            "Nº Rev.": text(entry.get("revisao_numero")) or "00",
            "Última": text(entry.get("ultima_revisao")),
            "Elaborado Por": text(entry.get("elaborador")),
            "Verificado Por": text(entry.get("verificador")),
            "Aprovado Por": text(entry.get("aprovador")),
            "Data Aprov.": text(entry.get("data_aprovacao")),
            "Localização em Diretório": text(entry.get("localizacao")),
            "Procedimento Raiz": text(entry.get("procedimento_raiz")),
            "Procedimentos Citados": "; ".join(as_list(entry.get("procedimentos_citados"))),
            "Observações": text(entry.get("observacoes")),
        }
        for name, value in values.items():
            ws.cell(row=row, column=cols[name], value=value)

        applied.append({"codigo": key, "row": row, "created": created, "codigo_titulo": codigo_titulo})

    # Renumber sequential Nº for filled rows.
    data_rows = iter_data_rows(ws, header_row, code_col)
    for idx, row in enumerate(data_rows, start=1):
        ws.cell(row=row, column=cols["Nº"], value=idx)

    return applied


def main() -> int:
    parser = argparse.ArgumentParser(description="Atualiza Lista Documental Mestra a partir do JSON v2")
    parser.add_argument("--content-json", required=True)
    parser.add_argument("--output-dir", required=True, help="Pasta documentos/ ou pasta de saída do processo")
    parser.add_argument("--template", default="")
    parser.add_argument("--lista-mestra-path", default="", help="Cópia de trabalho opcional para merge")
    parser.add_argument("--entries-json", default="", help="Caminho opcional para dump das entradas aplicadas")
    args = parser.parse_args()

    content_path = Path(args.content_json)
    output_dir = Path(args.output_dir)
    if not content_path.is_file():
        print(f"JSON não encontrado: {content_path}", file=sys.stderr)
        return 2

    script_dir = Path(__file__).resolve().parent
    default_template = script_dir.parent / "assets" / "templates" / "LISTA-MESTRA-template.xlsx"
    template_path = Path(args.template) if args.template else default_template
    ensure_template(template_path)

    output_dir.mkdir(parents=True, exist_ok=True)
    # Accept either .../documentos or process root.
    docs_dir = output_dir if output_dir.name.lower() == "documentos" else output_dir / "documentos"
    docs_dir.mkdir(parents=True, exist_ok=True)
    geracao_dir = docs_dir.parent / "geracao"
    geracao_dir.mkdir(parents=True, exist_ok=True)

    target = docs_dir / OUTPUT_NAME
    source_master = Path(args.lista_mestra_path) if args.lista_mestra_path else None
    if source_master and source_master.is_file():
        resolved_source = source_master.resolve()
        resolved_target = target.resolve() if target.exists() else target
        if resolved_source != resolved_target:
            # Never write back into referencias/; always work on the output copy.
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_master, target)
    elif not target.is_file():
        shutil.copy2(template_path, target)

    content = load_json(content_path)
    entries = build_entries_from_content(content)
    if not entries:
        print("Nenhuma entrada de lista mestra para aplicar.", file=sys.stderr)
        return 1

    wb = load_workbook(target)
    ws = find_sheet(wb)
    applied = upsert(ws, entries)
    wb.save(target)

    entries_path = Path(args.entries_json) if args.entries_json else geracao_dir / "lista-mestra-entradas.json"
    payload = {"arquivo": str(target), "entradas": entries, "aplicadas": applied}
    entries_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    created = sum(1 for item in applied if item["created"])
    updated = len(applied) - created
    print(f"Lista mestra: {target}")
    print(f"Entradas aplicadas: {len(applied)} (novas={created}, atualizadas={updated})")
    print(f"Dump: {entries_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
