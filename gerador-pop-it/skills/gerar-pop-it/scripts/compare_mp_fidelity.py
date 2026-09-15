#!/usr/bin/env python3
"""Compara a estrutura do MP gerado com a referência institucional MP.FIS.001."""

from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

import openpyxl


def _drawing_stats(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        xml = ""
        for name in archive.namelist():
            if name.startswith("xl/drawings/drawing") and name.endswith(".xml"):
                xml += archive.read(name).decode("utf-8")
    return {
        "shapes": xml.count("<sp ") + xml.count("<xdr:sp"),
        "connectors": xml.count("<cxnSp") + xml.count("<xdr:cxnSp"),
        "pictures": xml.count("<pic") + xml.count("<xdr:pic"),
        "ellipses": xml.count('prst="ellipse"'),
        "arrows": xml.count('tailEnd type="arrow"'),
    }


def compare(generated: Path, reference: Path) -> list[str]:
    failures: list[str] = []
    gen_wb = openpyxl.load_workbook(generated)
    ref_wb = openpyxl.load_workbook(reference)
    gen = gen_wb["MP"] if "MP" in gen_wb.sheetnames else gen_wb.worksheets[0]
    ref = ref_wb.worksheets[0]

    used_risk_rows = [
        row
        for row in range(40, 66)
        if not gen.row_dimensions[row].hidden
        and (
            gen.cell(row=row, column=2).value not in (None, "")
            or gen.cell(row=row, column=4).value not in (None, "")
            or gen.cell(row=row, column=12).value not in (None, "")
        )
    ]
    step_count = max(len(used_risk_rows), 1)
    risk_count = len(used_risk_rows) if used_risk_rows else 1

    checks = [
        ("aba-principal", gen.title == "MP", gen.title),
        ("zoom", gen.sheet_view.zoomScale == 85, gen.sheet_view.zoomScale),
        ("gridlines", gen.sheet_view.showGridLines is False, gen.sheet_view.showGridLines),
        ("freeze", str(gen.freeze_panes) == "A7", gen.freeze_panes),
        ("linhas", gen.max_row >= 86, gen.max_row),
        (
            "merges-base",
            len(gen.merged_cells.ranges) >= 165,
            len(gen.merged_cells.ranges),
        ),
        ("map-header-height", gen.row_dimensions[39].height == 120.0, gen.row_dimensions[39].height),
        ("risk-row-height", gen.row_dimensions[40].height == 119.25, gen.row_dimensions[40].height),
        ("inicio-height", gen.row_dimensions[5].height == 26.25, gen.row_dimensions[5].height),
        ("resultado-hidden", bool(gen.row_dimensions[6].hidden), gen.row_dimensions[6].hidden),
        ("paper", gen.page_setup.paperSize == 9, gen.page_setup.paperSize),
        (
            "margin-left",
            abs(float(gen.page_margins.left) - 0.3937007874015748) < 1e-6,
            gen.page_margins.left,
        ),
        ("print-area", "A$1:$J$86" in str(gen.print_area).replace("'", ""), gen.print_area),
        ("cf-count", len(list(gen.conditional_formatting)) == 2, len(list(gen.conditional_formatting))),
        ("depart-no-navy", gen["B3"].fill.fill_type in (None, "none"), gen["B3"].fill.fill_type),
        ("processo-no-navy", gen["D3"].fill.fill_type in (None, "none"), gen["D3"].fill.fill_type),
        (
            "inicio-navy",
            (gen["A5"].fill.fgColor.rgb if gen["A5"].fill.fill_type else None) in ("FF002060", "002060"),
            gen["A5"].fill.fgColor.rgb if gen["A5"].fill.fill_type else None,
        ),
        ("map-headers", gen["A39"].value == "ETAPAS", gen["A39"].value),
        ("widths-A", round(float(gen.column_dimensions["A"].width or 0), 2) == 23.43, gen.column_dimensions["A"].width),
    ]

    for name, ok, detail in checks:
        if not ok:
            failures.append(f"{name}: {detail}")

    # Ocultação estrutural (N riscos): linha 6 + SIPOC 16–37 + slots ociosos 40+N…65.
    # Baseline da referência FIS.001 (N=4) dá 45 ocultas; N>4 reduz o total sem ser defeito.
    sipoc_not_hidden = [row for row in range(16, 38) if not gen.row_dimensions[row].hidden]
    if sipoc_not_hidden:
        failures.append(f"hidden-sipoc: esperados ocultos 16-37, visiveis {sipoc_not_hidden}")

    used_start = 40
    used_end = 39 + risk_count
    idle_start = 40 + risk_count
    used_hidden = [
        row for row in range(used_start, used_end + 1) if gen.row_dimensions[row].hidden
    ]
    if used_hidden:
        failures.append(
            f"hidden-risk-used: esperados visiveis {used_start}-{used_end} (N={risk_count}), ocultos {used_hidden}"
        )
    idle_visible = [
        row for row in range(idle_start, 66) if not gen.row_dimensions[row].hidden
    ]
    if idle_visible:
        failures.append(
            f"hidden-risk-slots: esperados ocultos {idle_start}-65 (N={risk_count}), visiveis {idle_visible}"
        )

    # Larguras A–J (exceto G, que usa padrão).
    expected_widths = {
        "A": 23.43,
        "B": 20.57,
        "C": 32.86,
        "D": 29.0,
        "E": 30.86,
        "F": 5.71,
        "H": 6.86,
        "I": 22.43,
        "J": 25.86,
    }
    for column, expected in expected_widths.items():
        got = round(float(gen.column_dimensions[column].width or 0), 2)
        if got != expected:
            failures.append(f"width-{column}: {got} != {expected}")

    drawing = _drawing_stats(generated)
    expected_shapes = step_count
    expected_connectors = max(step_count - 1, 0)
    expected_ellipses = 1 if step_count == 1 else 2
    if drawing["pictures"] < 1:
        failures.append(f"logo-ausente: {drawing}")
    if drawing["shapes"] < expected_shapes:
        failures.append(f"formas-insuficientes: {drawing} esperado>={expected_shapes}")
    if drawing["connectors"] < expected_connectors:
        failures.append(f"conectores-ausentes: {drawing} esperado>={expected_connectors}")
    if drawing["ellipses"] < expected_ellipses:
        failures.append(f"elipses-insuficientes: {drawing} esperado>={expected_ellipses}")
    if expected_connectors and drawing["arrows"] < expected_connectors:
        failures.append(f"setas-ausentes: {drawing} esperado>={expected_connectors}")

    # Merges BARREIRA (E) e RESULTADO (J) no padrão da referência.
    first_used = used_risk_rows[0] if used_risk_rows else 40
    last_used = used_risk_rows[-1] if used_risk_rows else 40
    merge_set = {str(item) for item in gen.merged_cells.ranges}
    expected_barrier_merge = f"E{first_used}:E{last_used}"
    expected_result_merge = f"J{first_used}:J65"
    if first_used != last_used and expected_barrier_merge not in merge_set:
        failures.append(f"merge-barreira-ausente: esperado {expected_barrier_merge}")
    if expected_result_merge not in merge_set:
        failures.append(f"merge-resultado-ausente: esperado {expected_result_merge}")

    # Documentos e resultado uma vez só (topo do merge).
    barrier_top = gen.cell(row=first_used, column=5)
    result_top = gen.cell(row=first_used, column=10)
    if barrier_top.value in (None, ""):
        failures.append("barreira-vazia-no-topo-do-merge")
    else:
        for row in used_risk_rows[1:]:
            below = gen.cell(row=row, column=5).value
            if below not in (None, "", barrier_top.value):
                failures.append(f"barreira-repetida-{row}: {below!r}")
                break
    if result_top.value in (None, ""):
        # Fixture mínimo pode omitir resultado; exige só o merge e ausência de repetição.
        for row in used_risk_rows[1:]:
            below = gen.cell(row=row, column=10).value
            if below not in (None, ""):
                failures.append(f"resultado-repetido-{row}: {below!r}")
                break
    else:
        for row in used_risk_rows[1:]:
            below = gen.cell(row=row, column=10).value
            if below not in (None, "", result_top.value):
                failures.append(f"resultado-repetido-{row}: {below!r}")
                break

    for cell, label in (
        (barrier_top, "barreira"),
        (result_top, "resultado"),
        (gen.cell(row=first_used, column=9), "mitigacao"),
    ):
        if not getattr(cell.alignment, "wrap_text", False):
            failures.append(f"{label}-sem-wrap")

    # BARREIRA sem hiperlink azul (só no topo do merge).
    cell = barrier_top
    if cell.value not in (None, ""):
        color = None
        if cell.font.color is not None and cell.font.color.type == "rgb":
            color = cell.font.color.rgb
        if color in ("000563C1", "FF0563C1", "0563C1"):
            failures.append(f"barreira-hiperlink-{first_used}: {color}")
        if cell.font.underline not in (None, "none"):
            failures.append(f"barreira-underline-{first_used}: {cell.font.underline}")

    # Dropdown de classificação (lista Plan1!C1:C3), quem faz (Plan1!B) e DEPART (Plan1!D).
    list_formulas = []
    for dv in gen.data_validations.dataValidation:
        if dv.type == "list":
            list_formulas.append((str(dv.formula1 or ""), str(dv.sqref or "")))
    class_ok = any(
        "Plan1!$C$1:$C$3" in formula.replace("'", "") and "D40" in sqref
        for formula, sqref in list_formulas
    )
    who_ok = any(
        "Plan1!$B$1:$B$18" in formula.replace("'", "") and "B40" in sqref
        for formula, sqref in list_formulas
    )
    dept_ok = any(
        "Plan1!$D$1:$D$13" in formula.replace("'", "") and "C3" in sqref
        for formula, sqref in list_formulas
    )
    if not class_ok:
        failures.append(f"dropdown-classificacao-ausente: {list_formulas}")
    if not who_ok:
        failures.append(f"dropdown-quem-faz-ausente: {list_formulas}")
    if not dept_ok:
        failures.append(f"dropdown-departamento-ausente: {list_formulas}")

    if "Plan1" not in gen_wb.sheetnames:
        failures.append("plan1-ausente")
    else:
        plan1 = gen_wb["Plan1"]
        expected = ["TOLERÁVEL", "RAZOÁVEL (ALARP)", "INACEITÁVEL"]
        got = [plan1.cell(row=i, column=3).value for i in range(1, 4)]
        if got != expected:
            failures.append(f"plan1-classificacao: {got}")
        expected_depts = [
            "01 - Atendimento",
            "02 - Departamento Pessoal",
            "03 - Fiscal",
            "04 - Contábil",
            "05 - Paralegal",
            "06 - Financeiro",
            "07 - Comercial",
            "08 - Tecnologia da Informação",
            "09 - Processos e Qualidade",
            "10 - Sucesso do Cliente",
            "11 - Recursos Humanos",
            "12 - Marketing",
            "13 - Auditoria",
        ]
        got_depts = [plan1.cell(row=i, column=4).value for i in range(1, 14)]
        if got_depts != expected_depts:
            failures.append(f"plan1-departamentos: {got_depts}")

    # BARREIRA deve trazer o documento FORM (título), não só Bloco B0x.
    barrier_text = str(barrier_top.value or "")
    if barrier_text and "FORM." in barrier_text.upper():
        if "Bloco B0" in barrier_text and "FORM." in barrier_text.split("\n")[0]:
            first_line = barrier_text.split("\n")[0].strip()
            if "Bloco B0" in first_line:
                failures.append(f"barreira-form-bloco-em-vez-de-documento: {first_line!r}")
        # Exige linha FORM + pelo menos um de IN/PR quando o pacote os tiver no texto.
        lines = [ln.strip() for ln in barrier_text.splitlines() if ln.strip()]
        form_lines = [ln for ln in lines if ln.upper().startswith("FORM.")]
        if form_lines and "Bloco B0" in form_lines[0]:
            failures.append(f"barreira-form-nao-documento: {form_lines[0]!r}")

    if gen["E70"].value != "MUITO GRAVE":
        failures.append(f"legenda-muito-grave: {gen['E70'].value!r}")

    # Referência ainda precisa existir (ancora do template).
    if not reference.is_file():
        failures.append(f"referencia-ausente: {reference}")
    else:
        if len(ref.merged_cells.ranges) < 170:
            failures.append(f"referencia-merges-inesperados: {len(ref.merged_cells.ranges)}")

    return failures


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--generated", required=True)
    parser.add_argument("--reference", required=False)
    args = parser.parse_args()
    generated = Path(args.generated)
    if args.reference:
        reference = Path(args.reference)
    else:
        reference = next(Path("referencias").glob("01.3 MP.FIS.001*.xlsx"))
    failures = compare(generated, reference)
    if failures:
        print("FALHAS DE FIDELIDADE MP:")
        for item in failures:
            print(f"- {item}")
        return 1
    print("FIDELIDADE MP: OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
