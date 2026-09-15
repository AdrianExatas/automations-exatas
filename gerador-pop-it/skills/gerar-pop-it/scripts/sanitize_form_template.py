#!/usr/bin/env python3
"""Gera FORM-template.xlsx sanitizado a partir da referência FORM.QUA.002."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

ROOT = Path(__file__).resolve().parents[3]
REF = next((ROOT / "referencias").glob("FORM.QUA.002*.xlsx"))
OUT = Path(__file__).resolve().parents[1] / "assets" / "templates" / "FORM-template.xlsx"

# Blocos institucionais (header, coeficiente, itens, parecer).
FORM_BLOCKS = (
    {"header": 14, "coef": 15, "items": range(16, 18), "parecer": 18},
    {"header": 19, "coef": 20, "items": range(21, 23), "parecer": 23},
    {"header": 24, "coef": 25, "items": range(26, 28), "parecer": 28},
    {"header": 29, "coef": 30, "items": range(31, 33), "parecer": 33},
    {"header": 34, "coef": 35, "items": range(36, 46), "parecer": 46},
    {"header": 47, "coef": 48, "items": range(49, 64), "parecer": 64},
    {"header": 65, "coef": 66, "items": range(67, 79), "parecer": 79},
)

META_CLEAR = ("C3", "D3", "F2", "F4", "I4", "C9", "C12", "E12")
SIGNATURE_CELLS = ("B85", "F85", "I85")
SUMMARY_TITLE_CELLS = ("G9", "G10", "G11", "G12", "K9", "K10", "K11")


def clear_cell_value(ws, address: str) -> None:
    cell = ws[address]
    if isinstance(cell, MergedCell):
        return
    cell.value = None
    if cell.comment is not None:
        cell.comment = None


def set_cell_value(ws, address: str, value) -> None:
    cell = ws[address]
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def main() -> int:
    if not REF.is_file():
        raise SystemExit(f"Referência não encontrada: {REF}")

    # Import local para restaurar o DrawingML após o save do openpyxl.
    from build_excel_openpyxl import restore_form_drawing

    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REF, OUT)

    wb = openpyxl.load_workbook(OUT)
    ws = wb.worksheets[0]
    ws.title = "FORM"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    # Não usar freeze_panes=None via openpyxl — corrompe o pacote.
    # O freeze é removido em restore_form_drawing() via sheetViews.

    for addr in META_CLEAR:
        clear_cell_value(ws, addr)

    # Placeholders institucionais de contexto/cliente.
    set_cell_value(ws, "C12", "[CNPJ]")
    set_cell_value(ws, "E12", "[NOME CLIENTE]")
    set_cell_value(ws, "I4", "XX/XX/XXXX")
    set_cell_value(ws, "J4", "Versão: 00")
    set_cell_value(ws, "F2", "MODELO CHECK LIST TAREFA")

    for index, block in enumerate(FORM_BLOCKS, start=1):
        set_cell_value(ws, f"B{block['header']}", f"{index} | XXX")
        set_cell_value(ws, f"G{block['header']}", "SIM / NÃO")
        for row in block["items"]:
            clear_cell_value(ws, f"B{row}")
            clear_cell_value(ws, f"G{row}")
        set_cell_value(ws, f"F{block['parecer']}", "NÃO CONFORME")

    for addr in SUMMARY_TITLE_CELLS:
        # Mantém rótulos 1|XXX … 7|XXX no resumo.
        current = ws[addr].value
        if isinstance(current, str) and "|" in current:
            number = current.split("|", 1)[0].strip()
            set_cell_value(ws, addr, f"{number} | XXX")

    set_cell_value(ws, "B80", "OBSERVAÇÃO: \n")
    set_cell_value(ws, "B81", "Frequência: Sob Demanda")
    set_cell_value(ws, "B82", "Obs1: ")
    set_cell_value(ws, "B83", "Obs2: ")

    for addr in SIGNATURE_CELLS:
        clear_cell_value(ws, addr)

    # Remove nomes definidos de processo legado, se existirem.
    for name in list(wb.defined_names):
        lowered = name.lower()
        if lowered.startswith("form_") or lowered.startswith("processo"):
            del wb.defined_names[name]

    ws.print_area = "B2:K85"
    wb.save(OUT)
    restore_form_drawing(OUT, REF)
    print(f"Template sanitizado: {OUT}")
    print(f"Referência origem: {REF}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
