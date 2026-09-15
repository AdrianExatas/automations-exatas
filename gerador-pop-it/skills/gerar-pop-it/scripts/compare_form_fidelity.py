#!/usr/bin/env python3
"""Compara a estrutura do FORM gerado com a referência institucional FORM.QUA.002."""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

import openpyxl


def _drawing_stats(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        xml = ""
        has_xdr = False
        for name in archive.namelist():
            if name.startswith("xl/drawings/drawing") and name.endswith(".xml"):
                chunk = archive.read(name).decode("utf-8")
                xml += chunk
                if "xdr:wsDr" in chunk or 'xmlns:xdr=' in chunk:
                    has_xdr = True
            if name.startswith("xl/media/"):
                xml += name
    return {
        "pictures": xml.count("<pic") + xml.count("<xdr:pic"),
        "media": xml.count("xl/media/"),
        "xdr": has_xdr,
    }


def compare(generated: Path, reference: Path) -> list[str]:
    failures: list[str] = []
    gen_wb = openpyxl.load_workbook(generated)
    ref_wb = openpyxl.load_workbook(reference)
    gen = gen_wb["FORM"] if "FORM" in gen_wb.sheetnames else gen_wb.worksheets[0]
    ref = ref_wb.worksheets[0]

    checks = [
        ("aba-principal", gen.title == "FORM", gen.title),
        ("zoom", gen.sheet_view.zoomScale == 85, gen.sheet_view.zoomScale),
        ("gridlines", gen.sheet_view.showGridLines is False, gen.sheet_view.showGridLines),
        ("freeze", gen.freeze_panes in (None, "A1"), gen.freeze_panes),
        ("linhas", gen.max_row >= 85, gen.max_row),
        (
            "merges-base",
            len(gen.merged_cells.ranges) >= 165,
            len(gen.merged_cells.ranges),
        ),
        (
            "cf-count",
            len(list(gen.conditional_formatting)) >= 15,
            len(list(gen.conditional_formatting)),
        ),
        ("paper", gen.page_setup.paperSize == 9, gen.page_setup.paperSize),
        (
            "print-area",
            "B$2:$K$85" in str(gen.print_area).replace("'", "").replace("FORM!", ""),
            gen.print_area,
        ),
        ("coef-header", gen["B7"].value == "1 - Coeficiente de Conformidade", gen["B7"].value),
        ("resumo-header", gen["G7"].value == "Resumo de Inspeção", gen["G7"].value),
        ("observacao", str(gen["B80"].value or "").upper().startswith("OBSERVAÇÃO"), gen["B80"].value),
        ("assinatura-labels", gen["B84"].value == "Elaborado Por:", gen["B84"].value),
        (
            "widths-B",
            abs(float(gen.column_dimensions["B"].width or 0) - 12.42578125) < 0.2,
            gen.column_dimensions["B"].width,
        ),
    ]

    for name, ok, detail in checks:
        if not ok:
            failures.append(f"{name}: {detail}")

    # Dropdowns institucionais SIM/NÃO e CONFORME/NÃO CONFORME.
    list_formulas = []
    for dv in gen.data_validations.dataValidation:
        if dv.type == "list":
            list_formulas.append((str(dv.formula1 or ""), str(dv.sqref or "")))
    sim_ok = any("M$9" in formula.replace("'", "") or "$M$9" in formula for formula, _ in list_formulas)
    parecer_ok = any(
        "M$11" in formula.replace("'", "") or "$M$11" in formula for formula, _ in list_formulas
    )
    if not sim_ok:
        failures.append(f"dropdown-sim-nao-ausente: {list_formulas}")
    if not parecer_ok:
        failures.append(f"dropdown-parecer-ausente: {list_formulas}")

    drawing = _drawing_stats(generated)
    if drawing["pictures"] < 1:
        failures.append(f"logo-ausente: {drawing}")
    if not drawing["xdr"]:
        failures.append(f"drawing-sem-xdr: {drawing}")
    with zipfile.ZipFile(generated) as archive:
        rels_name = "xl/worksheets/_rels/sheet1.xml.rels"
        rels = ""
        if rels_name in archive.namelist():
            rels = archive.read(rels_name).decode("utf-8")
            if 'Target="/xl/drawings/' in rels:
                failures.append("drawing-target-absoluto")
            if "../drawings/drawing1.xml" not in rels:
                failures.append("drawing-target-relativo-ausente")
            if "../drawings/vmlDrawing1.vml" not in rels:
                failures.append("vml-target-relativo-ausente")
            if "../comments1.xml" not in rels:
                failures.append("comments-target-relativo-ausente")
            if "SEMED" in rels.upper() or "Licita" in rels:
                failures.append("hiperlink-legado-semed")
        wb_rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        if 'Target="/xl/' in wb_rels:
            failures.append("workbook-rels-target-absoluto")
        types_xml = archive.read("[Content_Types].xml").decode("utf-8")
        if "/xl/comments/comment1.xml" in types_xml:
            failures.append("content-types-comment1-quebrado")
        if 'PartName="/xl/comments1.xml"' not in types_xml:
            failures.append("content-types-comments1-ausente")
        if 'Extension="vml"' not in types_xml:
            failures.append("content-types-vml-ausente")
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        if "SEMED" in sheet_xml.upper() or "anysvml" in sheet_xml:
            failures.append("sheet-hiperlink-ou-vml-quebrado")
        if not re.search(r'<drawing\b[^>]*r:id="rId\d+"', sheet_xml):
            failures.append("sheet-drawing-rid-ausente")
        if not re.search(r'<legacyDrawing\b[^>]*r:id="rId\d+"', sheet_xml):
            failures.append("sheet-legacyDrawing-rid-ausente")
        if "xl/drawings/vmlDrawing1.vml" not in archive.namelist():
            failures.append("vml-ausente")
        if any(name.endswith("commentsDrawing1.vml") for name in archive.namelist()):
            failures.append("commentsDrawing1-residual")
        if "xl/comments1.xml" not in archive.namelist():
            failures.append("comments1-ausente")
        else:
            comments_xml = archive.read("xl/comments1.xml").decode("utf-8")
            for expected_ref in ("E8", "C9"):
                if f'ref="{expected_ref}"' not in comments_xml:
                    failures.append(f"comentario-{expected_ref}-ausente")

    # Cabeçalho preenchido (não placeholder do template).
    title = str(gen["F2"].value or "")
    if not title or title.upper() == "MODELO CHECK LIST TAREFA":
        failures.append(f"titulo-nao-preenchido: {title!r}")
    code = str(gen["D3"].value or "")
    if not code.upper().startswith("FORM."):
        failures.append(f"codigo-invalido: {code!r}")

    # Ao menos um bloco institucional com coeficiente parcial.
    block_headers = [
        str(gen[f"B{row}"].value or "")
        for row in (14, 19, 24, 29, 34, 47, 65)
        if not gen.row_dimensions[row].hidden
    ]
    if not any("|" in header for header in block_headers):
        failures.append(f"blocos-ausentes: {block_headers}")
    if not any(
        str(gen[f"B{row}"].value or "") == "Coeficiente Parcial:"
        for row in (15, 20, 25, 30, 35, 48, 66)
        if not gen.row_dimensions[row].hidden
    ):
        failures.append("coeficiente-parcial-ausente")

    # Fórmulas institucionais: resumo J9:J11 e faixas F{coef} do template.
    for addr in ("J9", "J10", "J11"):
        value = str(gen[addr].value or "")
        if not re.match(r"^=G\d+$", value):
            failures.append(f"resumo-{addr}-formula-ausente: {value!r}")
    for addr, expected in (
        ("F15", "=AVERAGE(M16:M17)"),
        ("F35", "=AVERAGE(M36:M45)"),
        ("F48", "=AVERAGE(M49:M63)"),
    ):
        value = str(gen[addr].value or "").replace(" ", "")
        if value.upper() != expected.upper():
            failures.append(f"coef-{addr}-faixa-inesperada: {gen[addr].value!r}")
    if gen["C9"].value in (None, ""):
        failures.append("c9-inicio-ausente")
    if str(gen["C11"].value or "").replace(" ", "").upper() != "=C10":
        failures.append(f"c11-inesperado: {gen['C11'].value!r}")
    if not str(gen["F7"].value or "").upper().startswith("=AVERAGE("):
        failures.append(f"f7-media-ausente: {gen['F7'].value!r}")
    if not str(gen["K12"].value or "").upper().startswith("=AVERAGE("):
        failures.append(f"k12-media-ausente: {gen['K12'].value!r}")

    named = set(gen_wb.defined_names)
    for expected in ("FORM_RESPOSTAS", "FORM_PARECERES", "FORM_COEFICIENTES"):
        if expected not in named:
            failures.append(f"named-range-ausente: {expected}")

    elaborador = str(gen["B85"].value or "").strip()
    if elaborador:
        color = None
        font = gen["B85"].font
        if font.color is not None and font.color.type == "rgb":
            color = str(font.color.rgb).upper()
        if color not in ("FFFF0000", "FF0000", "00FF0000"):
            failures.append(f"elaborado-nao-vermelho: {color}")

    if not reference.is_file():
        failures.append(f"referencia-ausente: {reference}")
    else:
        if len(ref.merged_cells.ranges) < 165:
            failures.append(f"referencia-merges-inesperados: {len(ref.merged_cells.ranges)}")
        if ref.sheet_view.zoomScale != 85:
            failures.append(f"referencia-zoom-inesperado: {ref.sheet_view.zoomScale}")

    return failures


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--generated", required=True)
    parser.add_argument("--reference", required=False)
    args = parser.parse_args()
    generated = Path(args.generated)
    if args.reference:
        reference = Path(args.reference)
    else:
        reference = next(Path("referencias").glob("FORM.QUA.002*.xlsx"))
    failures = compare(generated, reference)
    if failures:
        print("FALHAS DE FIDELIDADE FORM:")
        for item in failures:
            print(f"- {item}")
        return 1
    print("FIDELIDADE FORM: OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
