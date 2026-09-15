#!/usr/bin/env python3
"""Gera MP-template.xlsx sanitizado a partir da referência MP.FIS.001."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
REF = next((ROOT / "referencias").glob("01.3 MP.FIS.001*.xlsx"))
OUT = Path(__file__).resolve().parents[1] / "assets" / "templates" / "MP-template.xlsx"
LOGO = Path(__file__).resolve().parents[1] / "assets" / "exatas-logo.png"

# Faixas de dados do processo (valores limpos; estilos/estrutura preservados).
SIPOC_DATA_ROWS = range(10, 16)
RISK_ROWS = range(40, 66)
META_VALUE_CELLS = ("J1", "J2", "J3", "J4", "C3", "E3", "B5", "D5")
SIGNATURE_CELLS = ("A85", "D85", "G85")


def clear_cell_value(ws, address: str) -> None:
    cell = ws[address]
    # Células cobertas por merge são somente leitura.
    from openpyxl.cell.cell import MergedCell

    if isinstance(cell, MergedCell):
        return
    cell.value = None
    if cell.comment is not None:
        cell.comment = None


def main() -> int:
    if not REF.is_file():
        raise SystemExit(f"Referência não encontrada: {REF}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REF, OUT)

    wb = openpyxl.load_workbook(OUT)
    ws = wb.worksheets[0]
    ws.title = "MP"

    # Corrige typo institucional documentado no padrão.
    if isinstance(ws["A7"].value, str) and "FORNCEDOR" in ws["A7"].value:
        ws["A7"] = ws["A7"].value.replace("FORNCEDOR", "FORNECEDOR")
    ws["E70"] = "MUITO GRAVE"

    for addr in META_VALUE_CELLS:
        clear_cell_value(ws, addr)

    # Resultado esperado: mantém o rótulo, limpa o conteúdo do processo.
    ws["A6"] = "Resultado Esperado do Processo: "
    clear_cell_value(ws, "B5")  # preenchido na geração (=TODAY() ou emissão)

    for row in SIPOC_DATA_ROWS:
        for col in range(1, 11):
            from openpyxl.cell.cell import MergedCell
            from openpyxl.utils import get_column_letter

            addr = f"{get_column_letter(col)}{row}"
            cell = ws[addr]
            if isinstance(cell, MergedCell):
                continue
            if cell.value not in (None, ""):
                clear_cell_value(ws, addr)

    for row in RISK_ROWS:
        for col in (1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12):
            from openpyxl.utils import get_column_letter

            clear_cell_value(ws, f"{get_column_letter(col)}{row}")
        # Garante fórmula PXG em todos os slots.
        ws.cell(row=row, column=8).value = f"=F{row}*G{row}"
        # Slots extras ficam ocultos; a geração reabre conforme a quantidade.
        if row >= 44:
            ws.row_dimensions[row].hidden = True

    for addr in SIGNATURE_CELLS:
        clear_cell_value(ws, addr)

    # Remove named range hospitalar legado.
    for name in list(wb.defined_names):
        if name.lower().startswith("setores"):
            del wb.defined_names[name]

    wb.save(OUT)
    print(f"Template sanitizado: {OUT}")
    print(f"Referência origem: {REF}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
