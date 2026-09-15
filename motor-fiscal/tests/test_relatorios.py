"""M8 — relatorios XLSX fiscal + JSON tecnico, destino dossie e lote."""

import json
from pathlib import Path

from openpyxl import load_workbook

from motor_fiscal.__main__ import main
from motor_fiscal.orquestracao import listar_empresas_banco
from motor_fiscal.relatorios import (
    NOME_RESUMO_XLSX,
    SUBPASTA_TECNICO,
    gerar_relatorios,
    resolver_destino_relatorios,
)
from motor_fiscal.relatorios.destino import PASTA_09_DOSSIE
from motor_fiscal.relatorios.xlsx import LIMITE_PENDENCIAS
from tests.conftest import COMPETENCIA, EMPRESA, FIXTURES


def test_resolver_destino_padrao(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    dest = resolver_destino_relatorios(EMPRESA, COMPETENCIA)
    assert dest == tmp_path / "_local" / "relatorios" / EMPRESA / COMPETENCIA


def test_resolver_destino_dossie(tmp_path):
    dossie = tmp_path / EMPRESA / COMPETENCIA
    pasta09 = dossie / PASTA_09_DOSSIE
    pasta09.mkdir(parents=True)
    dest = resolver_destino_relatorios(EMPRESA, COMPETENCIA, dossie=dossie)
    assert dest == pasta09


def test_gerar_relatorios_xlsx_fiscal(tmp_path):
    resultado = {
        "empresa_cnpj": EMPRESA,
        "competencia": COMPETENCIA,
        "gerado_em": "2026-08-07T00:00:00+00:00",
        "modulos": {
            "documental": {
                "modulo": "documental",
                "empresa_cnpj": EMPRESA,
                "competencia": COMPETENCIA,
                "ok": False,
                "achados": [
                    {
                        "codigo": "DOCUMENTO_DUPLICADO",
                        "severidade": "erro",
                        "mensagem": "Chave 25260570120662004330550010000079201911578203 duplicada",
                        "detalhes": {
                            "chave": "25260570120662004330550010000079201911578203",
                            "ocorrencias": 2,
                        },
                    },
                    {
                        "codigo": "AVISO_X",
                        "severidade": "aviso",
                        "mensagem": "aviso de teste",
                    },
                ],
                "resumo": {"erros": 1, "avisos": 1, "total": 2},
            },
            "icms": {
                "modulo": "icms",
                "tributo": "icms",
                "ok": True,
                "achados": [],
                "resumo": {
                    "erros": 0,
                    "avisos": 0,
                    "total": 0,
                    "vl_icms_recolher": 100.0,
                    "cruzamentos_ok": 13,
                    "cruzamentos_total": 13,
                },
                "cruzamentos": [
                    {
                        "id": 13,
                        "nome": "EFD x guia",
                        "ok": True,
                        "soma_guia": 100.0,
                        "soma_efd": 100.0,
                    }
                ],
            },
        },
        "resumo": {
            "ok": False,
            "erros": 1,
            "avisos": 1,
            "modulos": {
                "documental": {"ok": False, "erros": 1, "avisos": 1, "total_achados": 2},
                "icms": {"ok": True, "erros": 0, "avisos": 0, "total_achados": 0},
            },
        },
    }
    dest = tmp_path / "out"
    # legado na raiz deve ser limpo
    dest.mkdir()
    (dest / "consolidado.json").write_text("{}", encoding="utf-8")
    (dest / "documental.json").write_text("{}", encoding="utf-8")

    pacote = gerar_relatorios(resultado, dest)
    assert Path(pacote["arquivos"]["consolidado_json"]) == dest / SUBPASTA_TECNICO / "consolidado.json"
    assert Path(pacote["arquivos"]["consolidado_json"]).is_file()
    assert not (dest / "consolidado.json").exists()
    assert not (dest / "documental.json").exists()

    resumo_xlsx = Path(pacote["arquivos"]["consolidado_xlsx"])
    assert resumo_xlsx == dest / NOME_RESUMO_XLSX
    assert resumo_xlsx.is_file()

    wb = load_workbook(resumo_xlsx)
    assert wb.sheetnames == ["Resumo", "Pendencias", "Por modulo"]
    ws = wb["Resumo"]
    campos = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, 12)}
    assert campos["Empresa (CNPJ)"] == EMPRESA
    assert campos["Resultado da conferencia"] == "Nao passou"
    assert campos["ICMS a recolher (E110)"] == 100.0
    assert campos["Soma das guias (cruzamento 13)"] == 100.0

    ws_pend = wb["Pendencias"]
    headers = [ws_pend.cell(row=1, column=c).value for c in range(1, 7)]
    assert headers == [
        "Modulo",
        "Gravidade",
        "Descricao",
        "Documento/Chave",
        "Valor",
        "Acao sugerida",
    ]
    assert ws_pend.cell(row=2, column=1).value == "Documental"
    assert ws_pend.cell(row=2, column=2).value == "Erro"

    ws_mod = wb["Por modulo"]
    assert ws_mod.cell(row=1, column=2).value == "Status"

    doc_xlsx = Path(pacote["arquivos"]["documental_xlsx"])
    assert doc_xlsx.name == "Documental.xlsx"
    wb_doc = load_workbook(doc_xlsx)
    assert "Resumo" in wb_doc.sheetnames
    assert "Pendencias" in wb_doc.sheetnames
    assert "Achados" not in wb_doc.sheetnames

    icms_xlsx = Path(pacote["arquivos"]["icms_xlsx"])
    assert icms_xlsx.name == "ICMS.xlsx"


def test_xlsx_limita_volume_pendencias(tmp_path):
    achados = [
        {
            "codigo": f"X{i}",
            "severidade": "erro",
            "mensagem": f"pendencia {i}",
            "detalhes": {"chave": f"{i:044d}"},
        }
        for i in range(LIMITE_PENDENCIAS + 25)
    ]
    resultado = {
        "empresa_cnpj": EMPRESA,
        "competencia": COMPETENCIA,
        "gerado_em": "2026-08-07T00:00:00+00:00",
        "modulos": {
            "documental": {
                "modulo": "documental",
                "ok": False,
                "achados": achados,
                "resumo": {"erros": len(achados), "avisos": 0, "total": len(achados)},
            }
        },
        "resumo": {
            "ok": False,
            "erros": len(achados),
            "avisos": 0,
            "modulos": {
                "documental": {
                    "ok": False,
                    "erros": len(achados),
                    "avisos": 0,
                    "total_achados": len(achados),
                }
            },
        },
    }
    pacote = gerar_relatorios(resultado, tmp_path / "vol")
    wb = load_workbook(pacote["arquivos"]["consolidado_xlsx"])
    ws = wb["Pendencias"]
    aviso = ws.cell(row=1, column=1).value
    assert "Exibindo" in str(aviso)
    assert str(LIMITE_PENDENCIAS) in str(aviso)
    # cabecalho na linha 2; dados a partir da 3 → no max LIMITE linhas de dados
    assert ws.cell(row=2, column=1).value == "Modulo"
    assert ws.cell(row=3 + LIMITE_PENDENCIAS, column=1).value is None


def test_cli_auditar_gera_xlsx_e_lote(tmp_path):
    db_dir = tmp_path / "db"
    assert main([
        "importar",
        "--empresa", EMPRESA,
        "--competencia", COMPETENCIA,
        "--xml-dir", str(FIXTURES / "xmls"),
        "--efd", str(FIXTURES / "efd_icms_ipi_2026-06.txt"),
        "--db-dir", str(db_dir),
    ]) == 0

    # segunda empresa sintetica (copia do banco)
    outro = "99888777000166"
    src = db_dir / EMPRESA / f"{COMPETENCIA}.db"
    dest_db = db_dir / outro
    dest_db.mkdir(parents=True)
    (dest_db / f"{COMPETENCIA}.db").write_bytes(src.read_bytes())

    assert listar_empresas_banco(COMPETENCIA, db_dir) == [EMPRESA, outro]

    saida_dir = tmp_path / "auditorias"
    rel_dir = tmp_path / "relatorios"
    codigo = main([
        "auditar",
        "--todas-empresas",
        "--competencia", COMPETENCIA,
        "--modulos", "documental",
        "--db-dir", str(db_dir),
        "--saida-dir", str(saida_dir),
        "--saida-relatorios", str(rel_dir),
    ])
    assert codigo in (0, 1)
    assert (saida_dir / f"lote-{COMPETENCIA}.json").is_file()
    assert (rel_dir / EMPRESA / COMPETENCIA / "Documental.xlsx").is_file()
    assert (rel_dir / EMPRESA / COMPETENCIA / NOME_RESUMO_XLSX).is_file()
    assert (rel_dir / outro / COMPETENCIA / SUBPASTA_TECNICO / "consolidado.json").is_file()


def test_cli_relatorio_dossie(tmp_path):
    db_dir = tmp_path / "db"
    assert main([
        "importar",
        "--empresa", EMPRESA,
        "--competencia", COMPETENCIA,
        "--efd", str(FIXTURES / "efd_icms_ipi_2026-06.txt"),
        "--db-dir", str(db_dir),
    ]) == 0
    dossie = tmp_path / "dossie" / EMPRESA / COMPETENCIA
    (dossie / PASTA_09_DOSSIE).mkdir(parents=True)
    codigo = main([
        "relatorio",
        "--empresa", EMPRESA,
        "--competencia", COMPETENCIA,
        "--modulos", "documental",
        "--db-dir", str(db_dir),
        "--dossie", str(dossie),
        "--saida-dir", str(tmp_path / "aud"),
        "--no-xlsx",  # deve ser ignorado com --dossie
    ])
    assert codigo == 0
    pasta09 = dossie / PASTA_09_DOSSIE
    assert (pasta09 / SUBPASTA_TECNICO / "documental.json").is_file()
    assert (pasta09 / NOME_RESUMO_XLSX).is_file()
    assert (pasta09 / "Documental.xlsx").is_file()
    assert not (pasta09 / "documental.json").exists()


def test_cli_relatorio_auditoria_json(tmp_path):
    aud = {
        "empresa_cnpj": EMPRESA,
        "competencia": COMPETENCIA,
        "gerado_em": "2026-08-07T00:00:00+00:00",
        "modulos": {
            "estoque": {
                "modulo": "estoque",
                "ok": True,
                "achados": [],
                "resumo": {"erros": 0, "avisos": 0, "total": 0},
            }
        },
        "resumo": {
            "ok": True,
            "erros": 0,
            "avisos": 0,
            "modulos": {"estoque": {"ok": True, "erros": 0, "avisos": 0, "total_achados": 0}},
        },
    }
    path_aud = tmp_path / "aud.json"
    path_aud.write_text(json.dumps(aud), encoding="utf-8")
    dossie = tmp_path / "dossie"
    (dossie / PASTA_09_DOSSIE).mkdir(parents=True)
    assert main([
        "relatorio",
        "--empresa", EMPRESA,
        "--competencia", COMPETENCIA,
        "--auditoria-json", str(path_aud),
        "--dossie", str(dossie),
    ]) == 0
    pasta09 = dossie / PASTA_09_DOSSIE
    assert (pasta09 / NOME_RESUMO_XLSX).is_file()
    assert (pasta09 / "Estoque.xlsx").is_file()
    assert (pasta09 / SUBPASTA_TECNICO / "consolidado.json").is_file()
