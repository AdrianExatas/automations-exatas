"""Geracao de planilhas XLSX fiscais (openpyxl) a partir de resumo + achados."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Limite por aba: Excel leve; detalhe completo fica no JSON tecnico.
LIMITE_PENDENCIAS = 1000

NOMES_MODULO: dict[str, str] = {
    "documental": "Documental",
    "icms": "ICMS",
    "ipi": "IPI",
    "pis_cofins": "PIS_COFINS",
    "estoque": "Estoque",
    "lmc": "LMC",
    "margens": "Margens",
}

_GRAVIDADE: dict[str, str] = {
    "erro": "Erro",
    "aviso": "Aviso",
    "info": "Info",
}

_ACOES: dict[str, str] = {
    "DOCUMENTO_DUPLICADO": "Remover duplicidade na origem ou na EFD",
    "DOCUMENTO_FALTANTE": "Localizar XML/CT-e e escriturar na EFD",
    "QUEBRA_SEQUENCIA": "Conferir numeracao e series faltantes",
    "CORRELACAO": "Conferir vinculo CT-e x NF-e / produto",
    "CRUZAMENTO": "Analisar divergencia do cruzamento e justificar ou corrigir",
    "MAPA_TRIBUTARIO": "Revisar CFOP/CST/aliquota na matriz da UF",
    "EFD_GUIA": "Conferir valor da guia x obrigacao E116/E110",
    "PIS_COFINS": "Conferir base/aliquota PIS/COFINS na EFD-Contribuicoes",
    "MARGEM": "Revisar margem / preco / CST do item",
}


def nome_modulo_exibicao(chave: str) -> str:
    return NOMES_MODULO.get(chave, chave.replace("_", " ").title())


def nome_arquivo_modulo(chave: str) -> str:
    return f"{nome_modulo_exibicao(chave)}.xlsx"


def re_sub_aba(nome: str) -> str:
    limpo = "".join(c if c.isalnum() or c in "-_" else "_" for c in nome)
    return (limpo or "aba")[:31]


def _gravidade(sev: Any) -> str:
    if sev is None:
        return ""
    return _GRAVIDADE.get(str(sev).lower(), str(sev))


def _acao_sugerida(codigo: Any, mensagem: str = "") -> str:
    cod = str(codigo or "").upper()
    for chave, acao in _ACOES.items():
        if chave in cod:
            return acao
    msg = (mensagem or "").lower()
    if "guia" in msg:
        return _ACOES["EFD_GUIA"]
    if "falt" in msg:
        return "Localizar documento e regularizar escrituracao"
    if "duplic" in msg:
        return _ACOES["DOCUMENTO_DUPLICADO"]
    if "sequen" in msg:
        return _ACOES["QUEBRA_SEQUENCIA"]
    return "Analisar pendencia e corrigir na origem ou na EFD"


def _doc_chave(detalhes: dict[str, Any], mensagem: str = "") -> str:
    for k in ("chave", "chave_nfe", "chave_cte", "documento", "num_doc", "n_doc"):
        if detalhes.get(k):
            return str(detalhes[k])
    # tenta extrair chave de 44 digitos da mensagem
    digitos = "".join(c for c in mensagem if c.isdigit())
    if len(digitos) >= 44:
        return digitos[:44]
    return ""


def _valor_detalhe(detalhes: dict[str, Any]) -> Any:
    for k in (
        "valor",
        "vl_or",
        "vl_opr",
        "vl_icms",
        "diferenca",
        "soma_guia",
        "soma_efd",
        "ocorrencias",
    ):
        if k in detalhes and detalhes[k] is not None:
            return detalhes[k]
    return ""


def _linha_pendencia(
    *,
    modulo: str,
    codigo: Any,
    severidade: Any,
    mensagem: str,
    detalhes: dict[str, Any] | None = None,
) -> dict[str, Any]:
    det = detalhes if isinstance(detalhes, dict) else {}
    return {
        "Modulo": nome_modulo_exibicao(modulo),
        "Gravidade": _gravidade(severidade),
        "Descricao": mensagem or str(codigo or ""),
        "Documento/Chave": _doc_chave(det, mensagem or ""),
        "Valor": _valor_detalhe(det),
        "Acao sugerida": _acao_sugerida(codigo, mensagem or ""),
    }


def _pendencias_de_achados(modulo: str, achados: list | None) -> list[dict[str, Any]]:
    saida: list[dict[str, Any]] = []
    for a in achados or []:
        if not isinstance(a, dict):
            continue
        saida.append(
            _linha_pendencia(
                modulo=modulo,
                codigo=a.get("codigo"),
                severidade=a.get("severidade"),
                mensagem=str(a.get("mensagem") or ""),
                detalhes=a.get("detalhes") if isinstance(a.get("detalhes"), dict) else {},
            )
        )
    return saida


def _pendencias_sinteticas_icms(payload: dict) -> list[dict[str, Any]]:
    """Quando o modulo ICMS nao popula achados, resume divergencias principais."""
    saida: list[dict[str, Any]] = []
    for cruz in payload.get("cruzamentos") or []:
        if not isinstance(cruz, dict) or cruz.get("ok", True):
            continue
        nome = cruz.get("nome") or f"Cruzamento {cruz.get('id', '')}"
        det = {
            "soma_efd": cruz.get("soma_efd"),
            "soma_guia": cruz.get("soma_guia"),
            "qtd_xml": cruz.get("qtd_xml"),
            "qtd_efd": cruz.get("qtd_efd"),
        }
        falt_efd = cruz.get("faltantes_na_efd") or []
        falt_xml = cruz.get("faltantes_no_xml") or []
        extra = ""
        if isinstance(falt_efd, list) and falt_efd:
            extra = f" ({len(falt_efd)} faltantes na EFD)"
        elif isinstance(falt_xml, list) and falt_xml:
            extra = f" ({len(falt_xml)} faltantes no XML)"
        saida.append(
            _linha_pendencia(
                modulo="icms",
                codigo="CRUZAMENTO",
                severidade="erro",
                mensagem=f"{nome} nao fecha{extra}",
                detalhes={k: v for k, v in det.items() if v is not None},
            )
        )
    mapa = payload.get("mapa_tributario") or {}
    for div in (mapa.get("divergencias") or [])[:50]:
        if not isinstance(div, dict):
            continue
        saida.append(
            _linha_pendencia(
                modulo="icms",
                codigo="MAPA_TRIBUTARIO",
                severidade="erro",
                mensagem=(
                    f"CFOP {div.get('cfop')} CST {div.get('cst_icms')}: "
                    f"aliquota {div.get('aliq_observada')} "
                    f"(esperada {div.get('aliq_esperada')})"
                ),
                detalhes={
                    "valor": div.get("vl_opr"),
                    "chave": "",
                },
            )
        )
    e116 = payload.get("e116") or {}
    if isinstance(e116, dict) and e116.get("ok") is False:
        saida.append(
            _linha_pendencia(
                modulo="icms",
                codigo="E116",
                severidade="aviso",
                mensagem=(
                    "E116: soma VL_OR diverge do ICMS a recolher do E110 "
                    f"(E110={e116.get('vl_icms_recolher_e110')}, "
                    f"soma OR={e116.get('soma_vl_or')})"
                ),
                detalhes={
                    "valor": e116.get("soma_vl_or"),
                },
            )
        )
    return saida


def coletar_pendencias_modulo(nome: str, payload: dict) -> list[dict[str, Any]]:
    pend = _pendencias_de_achados(nome, payload.get("achados"))
    if not pend and nome == "icms":
        pend = _pendencias_sinteticas_icms(payload)
    return pend


def coletar_pendencias_resultado(resultado: dict) -> list[dict[str, Any]]:
    # Erros primeiro, depois avisos; estavel por modulo
    ordem_sev = {"Erro": 0, "Aviso": 1, "Info": 2, "": 3}
    todas: list[dict[str, Any]] = []
    for nome, payload in (resultado.get("modulos") or {}).items():
        if not isinstance(payload, dict):
            continue
        todas.extend(coletar_pendencias_modulo(nome, payload))
    todas.sort(key=lambda r: (ordem_sev.get(str(r.get("Gravidade")), 9), str(r.get("Modulo"))))
    return todas


def _valores_icms_guia(resultado: dict) -> tuple[Any, Any, Any]:
    """Retorna (icms_a_recolher, soma_guia, cruzamento_13_ok)."""
    icms = (resultado.get("modulos") or {}).get("icms") or {}
    if not isinstance(icms, dict):
        return None, None, None
    resumo = icms.get("resumo") if isinstance(icms.get("resumo"), dict) else {}
    vl_recolher = resumo.get("vl_icms_recolher")
    soma_guia = None
    c13_ok = None
    for cruz in icms.get("cruzamentos") or []:
        if not isinstance(cruz, dict):
            continue
        if cruz.get("id") == 13 or "guia" in str(cruz.get("nome") or "").lower():
            soma_guia = cruz.get("soma_guia")
            c13_ok = cruz.get("ok")
            break
    return vl_recolher, soma_guia, c13_ok


def _escrever_pares(ws, pares: list[tuple[str, Any]], *, cabecalhos: tuple[str, str] = ("Campo", "Valor")) -> None:
    ws.cell(row=1, column=1, value=cabecalhos[0]).font = Font(bold=True)
    ws.cell(row=1, column=2, value=cabecalhos[1]).font = Font(bold=True)
    for i, (campo, valor) in enumerate(pares, start=2):
        ws.cell(row=i, column=1, value=campo)
        ws.cell(row=i, column=2, value="" if valor is None else valor)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 56


def _escrever_tabela(
    ws,
    linhas: list[dict[str, Any]],
    colunas: list[str] | None = None,
    *,
    total: int | None = None,
    limite: int | None = None,
) -> None:
    if not linhas:
        ws.cell(row=1, column=1, value="(sem pendencias)")
        return
    cols = colunas or list(linhas[0].keys())
    lim = limite if limite is not None else LIMITE_PENDENCIAS
    exibidas = linhas[:lim]
    tot = total if total is not None else len(linhas)
    row0 = 1
    if tot > len(exibidas):
        ws.cell(
            row=1,
            column=1,
            value=f"Exibindo {len(exibidas)} de {tot} pendencias (detalhe completo no JSON tecnico).",
        ).font = Font(italic=True, size=10)
        row0 = 2
    for col, nome in enumerate(cols, start=1):
        cell = ws.cell(row=row0, column=col, value=nome)
        cell.font = Font(bold=True)
    for i, item in enumerate(exibidas):
        for col, nome in enumerate(cols, start=1):
            valor = item.get(nome)
            if isinstance(valor, (dict, list)):
                valor = str(valor)[:500]
            ws.cell(row=row0 + 1 + i, column=col, value=valor)
    for col, nome in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(12, len(nome) + 4))


_COLS_PEND = ["Modulo", "Gravidade", "Descricao", "Documento/Chave", "Valor", "Acao sugerida"]
_COLS_PEND_MOD = ["Gravidade", "Descricao", "Documento/Chave", "Valor", "Acao sugerida"]


def escrever_xlsx_modulo(payload: dict, caminho: str | Path) -> Path:
    """XLSX fiscal do modulo: abas Resumo + Pendencias."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    nome = str(payload.get("modulo") or caminho.stem)
    resumo = payload.get("resumo") if isinstance(payload.get("resumo"), dict) else {}
    ok = payload.get("ok")
    if ok is None:
        ok = resumo.get("ok")
    erros = resumo.get("erros", 0)
    avisos = resumo.get("avisos", 0)

    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    pares: list[tuple[str, Any]] = [
        ("Modulo", nome_modulo_exibicao(nome)),
        ("Empresa (CNPJ)", payload.get("empresa_cnpj")),
        ("Competencia", payload.get("competencia")),
        ("Status", "OK" if ok else "Divergente"),
        ("Erros", erros),
        ("Avisos", avisos),
        ("Total de achados", resumo.get("total", len(payload.get("achados") or []))),
    ]
    if nome == "icms" or payload.get("tributo") == "icms":
        pares.append(("ICMS a recolher (E110)", resumo.get("vl_icms_recolher")))
        pares.append(("Saldo credor a transportar", resumo.get("vl_sld_credor_transportar")))
        pares.append(
            (
                "Cruzamentos OK",
                f"{resumo.get('cruzamentos_ok', '')}/{resumo.get('cruzamentos_total', '')}",
            )
        )
    _escrever_pares(ws_resumo, pares)

    pend = coletar_pendencias_modulo(nome, payload)
    ws_pend = wb.create_sheet("Pendencias")
    # sem coluna Modulo na planilha do proprio modulo
    linhas_mod = [{k: r[k] for k in _COLS_PEND_MOD} for r in pend]
    _escrever_tabela(ws_pend, linhas_mod, _COLS_PEND_MOD, total=len(pend))

    wb.save(caminho)
    return caminho


def escrever_xlsx_consolidado(resultado: dict, caminho: str | Path) -> Path:
    """XLSX fiscal consolidado: Resumo + Pendencias + Por modulo."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    resumo = resultado.get("resumo") if isinstance(resultado.get("resumo"), dict) else {}
    vl_recolher, soma_guia, c13_ok = _valores_icms_guia(resultado)

    ws = wb.active
    ws.title = "Resumo"
    status = "Passou" if resumo.get("ok") else "Nao passou"
    pares: list[tuple[str, Any]] = [
        ("Empresa (CNPJ)", resultado.get("empresa_cnpj")),
        ("Competencia", resultado.get("competencia")),
        ("Gerado em", resultado.get("gerado_em")),
        ("Resultado da conferencia", status),
        ("Total de erros", resumo.get("erros")),
        ("Total de avisos", resumo.get("avisos")),
        ("ICMS a recolher (E110)", vl_recolher),
        ("Soma das guias (cruzamento 13)", soma_guia),
        (
            "EFD x guia",
            "OK" if c13_ok is True else ("Divergente" if c13_ok is False else "N/A"),
        ),
    ]
    _escrever_pares(ws, pares)

    pend = coletar_pendencias_resultado(resultado)
    ws_pend = wb.create_sheet("Pendencias")
    _escrever_tabela(ws_pend, pend, _COLS_PEND, total=len(pend))

    mods_resumo = resumo.get("modulos") or {}
    linhas_mod: list[dict[str, Any]] = []
    for nome, info in mods_resumo.items():
        if not isinstance(info, dict):
            info = {}
        ok_mod = info.get("ok")
        linhas_mod.append(
            {
                "Modulo": nome_modulo_exibicao(nome),
                "Status": "OK" if ok_mod else "Divergente",
                "Erros": info.get("erros"),
                "Avisos": info.get("avisos"),
                "Total pendencias": info.get("total_achados"),
            }
        )
    # inclui modulos presentes no payload e ausentes do resumo
    for nome in (resultado.get("modulos") or {}):
        if nome not in mods_resumo:
            payload = resultado["modulos"][nome]
            ok_mod = payload.get("ok") if isinstance(payload, dict) else None
            linhas_mod.append(
                {
                    "Modulo": nome_modulo_exibicao(nome),
                    "Status": "OK" if ok_mod else "Divergente",
                    "Erros": "",
                    "Avisos": "",
                    "Total pendencias": len((payload or {}).get("achados") or [])
                    if isinstance(payload, dict)
                    else "",
                }
            )
    ws_mod = wb.create_sheet("Por modulo")
    _escrever_tabela(
        ws_mod,
        linhas_mod,
        ["Modulo", "Status", "Erros", "Avisos", "Total pendencias"],
        limite=len(linhas_mod) + 1,
    )

    wb.save(caminho)
    return caminho
