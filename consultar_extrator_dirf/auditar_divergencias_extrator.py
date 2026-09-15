import argparse
import os
import sys
from decimal import Decimal
from datetime import datetime, date
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_env, get_connection

def auditar_extrator(codi_emp=None, competencia=None, output_path="Divergencias_Extrator_DIRF.xlsx"):
    conn = get_connection()
    cursor = conn.cursor()

    print("=" * 90)
    print("AUDITORIA INTEGRADA: EXTRATOR DA DIRF x eSOCIAL (BANCO DE DADOS DOMÍNIO)")
    print("=" * 90)

    # 1. Obter dados do Extrator S-5002
    where_ext = []
    if codi_emp:
        where_ext.append(f"tot.CODI_EMP = {codi_emp}")
    if competencia:
        where_ext.append(f"tot.COMPETENCIA_REF = '{competencia}'")
    
    where_ext_str = ("WHERE " + " AND ".join(where_ext)) if where_ext else ""

    print("Consultando dados consolidados do Extrator S-5002...")
    query_extrator = f"""
        SELECT 
            tot.CODI_EMP,
            emp.nome_emp,
            tot.COMPETENCIA_REF,
            tot.CRMEN,
            ext.CPF,
            COALESCE(tot.VLRRENDTRIB, 0) as esocial_rendtrib,
            COALESCE(tot.VLRRENDTRIB13, 0) as esocial_rendtrib13,
            COALESCE(tot.VLRPREVOFICIAL, 0) as esocial_prevoficial,
            COALESCE(tot.VLRPREVOFICIAL13, 0) as esocial_prevoficial13,
            COALESCE(tot.VLRCRMEN, 0) as esocial_crmen,
            COALESCE(tot.VLRCR13MEN, 0) as esocial_cr13men,
            COALESCE(tot.VLRINDRESCONTRATO, 0) as esocial_indres,
            COALESCE(tot.VLRABONOPEC, 0) as esocial_abonopec,
            COALESCE(tot.VLRDIARIAS, 0) as esocial_diarias,
            COALESCE(tot.VLRAJUDACUSTO, 0) as esocial_ajudacusto,
            COALESCE(tot.VLRISENOUTROS, 0) as esocial_outros_isentos,
            COALESCE(tot.VLRAUXMORADIA, 0) as esocial_auxmoradia,
            tot.I_DADOS_EVENTOS
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_TOTALIZADOR tot
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = tot.CODI_EMP AND ext.I_DADOS_EVENTOS = tot.I_DADOS_EVENTOS
        JOIN bethadba.GEEMPRE emp ON emp.CODI_EMP = tot.CODI_EMP
        {where_ext_str}
    """
    cursor.execute(query_extrator)
    cols_ext = [desc[0].lower() for desc in cursor.description]
    df_ext = pd.DataFrame.from_records(cursor.fetchall(), columns=cols_ext)
    print(f"Total registros no extrator S-5002: {len(df_ext)}")

    if df_ext.empty:
        print("Nenhum dado de extrator encontrado para os filtros informados.")
        conn.close()
        return

    # 2. Consultar planos de saúde do extrator
    print("Consultando informações de planos de saúde do extrator...")
    query_saude = f"""
        SELECT 
            ps.CODI_EMP,
            ps.COMPETENCIA_REF,
            ext.CPF,
            SUM(COALESCE(ps.VLRSAUDETIT, 0)) as esocial_plano_saude
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_PLANO_SAUDE ps
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = ps.CODI_EMP AND ext.I_DADOS_EVENTOS = ps.I_DADOS_EVENTOS
        {where_ext_str.replace('tot.', 'ps.')}
        GROUP BY ps.CODI_EMP, ps.COMPETENCIA_REF, ext.CPF
    """
    cursor.execute(query_saude)
    cols_saude = [desc[0].lower() for desc in cursor.description]
    df_saude = pd.DataFrame.from_records(cursor.fetchall(), columns=cols_saude)

    # 3. Consultar movimentação da Folha (Sistema) das empresas envolvidas
    empresas_alvo = list(df_ext['codi_emp'].unique())
    print(f"Consultando folha de pagamento para {len(empresas_alvo)} empresa(s)...")

    emp_filter = f"m.codi_emp IN ({','.join(map(str, empresas_alvo))})"
    comp_alvo = [d.strftime('%Y-%m-%d') if isinstance(d, (datetime, date)) else str(d) for d in df_ext['competencia_ref'].unique()]
    comp_filter = f"m.data IN ({','.join([repr(c) for c in comp_alvo])})"

    query_mov = f"""
        SELECT 
            m.CODI_EMP,
            m.DATA as competencia,
            f.CPF,
            f.NOME as nome_colaborador,
            m.I_EVENTOS,
            ev.NOME as nome_evento,
            m.VALOR_CAL,
            ev.REND_TRIBUTAVEIS,
            ev.REND_ISENTOS,
            ev.REND_ISENTOS_DESCRICAO,
            ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL
        FROM bethadba.FOMOVTO m
        JOIN bethadba.FOEMPREGADOS f 
          ON f.CODI_EMP = m.CODI_EMP AND f.I_EMPREGADOS = m.I_EMPREGADOS
        JOIN bethadba.FOEVENTOS ev 
          ON ev.CODI_EMP = m.CODI_EMP AND ev.I_EVENTOS = m.I_EVENTOS
        WHERE {emp_filter} AND {comp_filter}
    """
    cursor.execute(query_mov)
    cols_mov = [desc[0].lower() for desc in cursor.description]
    df_mov = pd.DataFrame.from_records(cursor.fetchall(), columns=cols_mov)
    print(f"Total registros na folha (fomovto): {len(df_mov)}")

    # 4. Auditoria de Rubricas com Inconsistência Cadastral
    print("Analisando inconsistências no cadastro de rubricas (FOEVENTOS)...")
    query_rubricas = f"""
        SELECT 
            ev.CODI_EMP,
            emp.nome_emp,
            ev.I_EVENTOS,
            ev.NOME as nome_rubrica,
            ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL as inc_esocial,
            ev.REND_TRIBUTAVEIS as rend_tributaveis,
            ev.REND_ISENTOS as rend_isentos,
            ev.REND_ISENTOS_DESCRICAO as desc_isentos,
            CASE 
                WHEN ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (70, 71, 72, 73, 74, 75, 76, 77, 78, 79) AND ev.REND_ISENTOS = 0 
                THEN 'eSocial configurado como Isento (cód ' || CAST(ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL AS VARCHAR) || '), mas no Domínio NÃO está marcado como Isento'
                WHEN ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (11, 12, 13, 14, 15) AND ev.REND_TRIBUTAVEIS = 0 
                THEN 'eSocial configurado como Tributável (cód ' || CAST(ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL AS VARCHAR) || '), mas no Domínio NÃO está marcado como Tributável'
                ELSE 'Outro conflito'
            END as inconsistencia,
            CASE 
                WHEN ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (70, 71, 72, 73, 74, 75, 76, 77, 78, 79) AND ev.REND_ISENTOS = 0 
                THEN 'Na aba DIRF da Rubrica, configurar Rendimento Isento e Não Tributável (Ex: Outros) OU corrigir incidência no eSocial se for verba salarial'
                WHEN ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (11, 12, 13, 14, 15) AND ev.REND_TRIBUTAVEIS = 0 
                THEN 'Na aba DIRF da Rubrica, marcar como Rendimento Tributável OU corrigir incidência no eSocial'
                ELSE 'Revisar configuração da rubrica'
            END as recomendacao
        FROM bethadba.FOEVENTOS ev
        JOIN bethadba.GEEMPRE emp ON emp.CODI_EMP = ev.CODI_EMP
        WHERE ev.CODI_EMP IN ({','.join(map(str, empresas_alvo))})
          AND (
              (ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (70, 71, 72, 73, 74, 75, 76, 77, 78, 79) AND ev.REND_ISENTOS = 0)
              OR 
              (ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (11, 12, 13, 14, 15) AND ev.REND_TRIBUTAVEIS = 0)
          )
        ORDER BY ev.CODI_EMP, ev.I_EVENTOS
    """
    cursor.execute(query_rubricas)
    cols_rub = [desc[0].lower() for desc in cursor.description]
    df_rubricas = pd.DataFrame.from_records(cursor.fetchall(), columns=cols_rub)
    print(f"Total rubricas com inconsistência cadastral nessas empresas: {len(df_rubricas)}")

    # 5. Cruzamento Detalhado por Colaborador
    print("Processando conciliação por colaborador e gerando métricas...")
    divergencias_colaboradores = []
    
    for (codi, comp, cpf), grp_ext in df_ext.groupby(['codi_emp', 'competencia_ref', 'cpf']):
        nome_emp = grp_ext['nome_emp'].iloc[0]
        ext_outros = float(grp_ext['esocial_outros_isentos'].sum())
        ext_rendtrib = float(grp_ext['esocial_rendtrib'].sum())
        ext_rendtrib13 = float(grp_ext['esocial_rendtrib13'].sum())
        ext_prev = float(grp_ext['esocial_prevoficial'].sum()) + float(grp_ext['esocial_prevoficial13'].sum())
        ext_cr = float(grp_ext['esocial_crmen'].sum()) + float(grp_ext['esocial_cr13men'].sum())
        ext_indres = float(grp_ext['esocial_indres'].sum())
        ext_diarias = float(grp_ext['esocial_diarias'].sum())
        ext_ajuda = float(grp_ext['esocial_ajudacusto'].sum())

        # Plano de saúde extrator
        saude_cpf = df_saude[(df_saude['codi_emp'] == codi) & (df_saude['cpf'] == cpf)]
        ext_saude = float(saude_cpf['esocial_plano_saude'].sum()) if not saude_cpf.empty else 0.0

        # Movimentos da Folha
        grp_mov = df_mov[(df_mov['codi_emp'] == codi) & (df_mov['cpf'] == cpf)]
        nome_trab = grp_mov['nome_colaborador'].iloc[0] if not grp_mov.empty else "NÃO LOCALIZADO NA FOLHA"

        sis_outros = 0.0
        sis_rendtrib = 0.0
        sis_prev = 0.0
        sis_irrf = 0.0
        sis_indres = 0.0
        sis_saude = 0.0
        
        rubricas_divergentes = []

        for _, row in grp_mov.iterrows():
            vlr = float(row['valor_cal'])
            inc_esocial = row['codigo_incidencia_irrf_esocial']
            rend_trib = row['rend_tributaveis']
            rend_isen = row['rend_isentos']
            cod_eve = row['i_eventos']
            nom_eve = row['nome_evento']

            # Verifica conflito específico
            if inc_esocial == 79 and rend_isen == 0:
                rubricas_divergentes.append(f"Rubrica {cod_eve} ({nom_eve}): R$ {vlr:.2f} [eSocial cód 79 x Sistema rend_isentos=0]")
            elif inc_esocial in (11, 12, 13) and rend_trib == 0:
                rubricas_divergentes.append(f"Rubrica {cod_eve} ({nom_eve}): R$ {vlr:.2f} [eSocial cód 11 x Sistema rend_tributaveis=0]")

            if rend_isen == 7: # Outros isentos
                sis_outros += vlr
            elif rend_isen == 4: # Indenizações
                sis_indres += vlr
            
            if rend_trib == 1:
                sis_rendtrib += vlr
            elif rend_trib == 2: # INSS
                sis_prev += vlr
            elif rend_trib == 3: # IRRF
                sis_irrf += vlr

            # Detecção de plano de saúde na folha
            if 'plano de sa' in nom_eve.lower():
                sis_saude += vlr

        dif_outros = round(sis_outros - ext_outros, 2)
        dif_rendtrib = round(sis_rendtrib - ext_rendtrib, 2)
        dif_prev = round(sis_prev - ext_prev, 2)
        dif_irrf = round(sis_irrf - ext_cr, 2)

        tem_divergencia = (
            abs(dif_outros) > 0.01 or 
            abs(dif_rendtrib) > 0.01 or 
            abs(dif_prev) > 0.01 or 
            abs(dif_irrf) > 0.01 or 
            len(rubricas_divergentes) > 0
        )

        if tem_divergencia:
            divergencias_colaboradores.append({
                "codi_emp": codi,
                "nome_emp": nome_emp,
                "competencia": comp.strftime('%Y-%m') if isinstance(comp, (datetime, date)) else str(comp),
                "cpf": cpf,
                "nome_colaborador": nome_trab,
                "sis_outros_isentos": sis_outros,
                "esocial_outros_isentos": ext_outros,
                "dif_outros_isentos": dif_outros,
                "sis_rend_tributavel": sis_rendtrib,
                "esocial_rend_tributavel": ext_rendtrib,
                "dif_rend_tributavel": dif_rendtrib,
                "sis_prev_oficial": sis_prev,
                "esocial_prev_oficial": ext_prev,
                "dif_prev_oficial": dif_prev,
                "sis_imposto_retido": sis_irrf,
                "esocial_imposto_retido": ext_cr,
                "dif_imposto_retido": dif_irrf,
                "motivo_divergencia": "; ".join(rubricas_divergentes) if rubricas_divergentes else "Divergência de valores"
            })

    df_res_colab = pd.DataFrame(divergencias_colaboradores)

    # 6. Agrupamento por Empresa
    resumo_empresas = pd.DataFrame()
    if not df_res_colab.empty:
        resumo_empresas = df_res_colab.groupby(['codi_emp', 'nome_emp', 'competencia']).agg(
            qtd_colaboradores_divergentes=('cpf', 'count'),
            total_dif_outros_isentos=('dif_outros_isentos', 'sum'),
            total_dif_rend_tributavel=('dif_rend_tributavel', 'sum'),
            total_dif_prev_oficial=('dif_prev_oficial', 'sum'),
            total_dif_imposto_retido=('dif_imposto_retido', 'sum')
        ).reset_index()

    # 7. Formatação e Exportação Excel
    print(f"\nFormatando e exportando para Excel ({output_path})...")
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        if not resumo_empresas.empty:
            resumo_empresas.to_excel(writer, sheet_name="Resumo Empresas", index=False)
        else:
            pd.DataFrame([{"Mensagem": "Nenhuma divergência encontrada nas empresas selecionadas"}]).to_excel(writer, sheet_name="Resumo Empresas", index=False)
        
        if not df_res_colab.empty:
            df_res_colab.to_excel(writer, sheet_name="Detalhamento Colaboradores", index=False)
        
        if not df_rubricas.empty:
            df_rubricas.to_excel(writer, sheet_name="Rubricas Conflitantes", index=False)

    # Estilização visual com openpyxl
    import openpyxl
    wb = openpyxl.load_workbook(output_path)
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        
        # Formatar cabeçalho
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        # Formatar linhas de dados
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                
                # Formatação de valores numéricos
                col_name = ws.cell(row=1, column=cell.column).value or ""
                if any(k in str(col_name).lower() for k in ['dif_', 'sis_', 'esocial_', 'total_dif_']):
                    cell.number_format = 'R$ #,##0.00;[Red]R$ -#,##0.00;"R$ 0.00"'
                elif 'cpf' in str(col_name).lower():
                    cell.number_format = '@'

        # Ajustar largura das colunas automaticamente
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 65)

    wb.save(output_path)
    conn.close()

    print("=" * 90)
    print(f"AUDITORIA CONCLUÍDA COM SUCESSO!")
    print(f"Relatório salvo em: {os.path.abspath(output_path)}")
    print(f"Empresas analisadas: {len(empresas_alvo)}")
    print(f"Empresas com divergência: {len(resumo_empresas)}")
    print(f"Colaboradores com inconsistência: {len(df_res_colab)}")
    print(f"Rubricas com conflito cadastral: {len(df_rubricas)}")
    print("=" * 90)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Auditor de Divergências Extrator DIRF x eSocial")
    parser.add_argument("--empresa", type=int, help="Código da empresa no Domínio")
    parser.add_argument("--competencia", type=str, help="Competência (ex: 2026-06-01)")
    parser.add_argument("--output", type=str, default="Divergencias_Extrator_DIRF.xlsx", help="Caminho do arquivo de saída")
    args = parser.parse_args()
    auditar_extrator(args.empresa, args.competencia, args.output)
