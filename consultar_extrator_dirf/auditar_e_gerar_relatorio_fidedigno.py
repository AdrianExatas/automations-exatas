"""
=============================================================================
AUDITORIA E CONCILIAÇÃO EXECUTIVA: EXTRATOR DA DIRF x eSOCIAL (DOMÍNIO)
=============================================================================
Este script implementa o motor auditado de conciliação fiscal e gera
um relatório executivo de alta legibilidade para contadores e analistas de DP.
=============================================================================
"""

import sys
import os
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_env, get_connection, format_cpf, format_cnpj


def extrair_e_auditar_dados(ano=2026, codi_emp=None):
    print("=" * 85)
    print(f"AUDITORIA EXTRATOR DIRF x eSOCIAL - ANO BASE: {ano}")
    print("=" * 85)
    conn = get_connection()
    cur = conn.cursor()

    emp_filter_sql = f"AND tot.CODI_EMP = {codi_emp}" if codi_emp else ""

    # 1. eSocial S-5002 Totalizador Validado (Deduplicado por CPF e sem joins de expansão)
    print("\n[1/4] Extraindo dados do eSocial S-5002 Extrator Totalizador...")
    sql_es = f"""
        SELECT 
            tot.CODI_EMP,
            emp.NOME_EMP,
            emp.CGCE_EMP,
            tot.COMPETENCIA_REF,
            ext.CPF,
            COALESCE(emp_col.NOME, 'COLABORADOR CPF ' + ext.CPF) AS NOME_COLABORADOR,
            COALESCE(tot.VLRRENDTRIB, 0)        AS ESOCIAL_RENDTRIB_MENSAL,
            COALESCE(tot.VLRRENDTRIB13, 0)      AS ESOCIAL_RENDTRIB_13,
            COALESCE(tot.VLRPREVOFICIAL, 0)     AS ESOCIAL_PREV_MENSAL,
            COALESCE(tot.VLRPREVOFICIAL13, 0)   AS ESOCIAL_PREV_13,
            COALESCE(tot.VLRCRMEN, 0)           AS ESOCIAL_IRRF_MENSAL,
            COALESCE(tot.VLRCR13MEN, 0)         AS ESOCIAL_IRRF_13,
            COALESCE(tot.VLRPARCISENTA65, 0) + COALESCE(tot.VLRPARCISENTA65DEC, 0) +
            COALESCE(tot.VLRDIARIAS, 0) + COALESCE(tot.VLRAJUDACUSTO, 0) +
            COALESCE(tot.VLRINDRESCONTRATO, 0) + COALESCE(tot.VLRABONOPEC, 0) +
            COALESCE(tot.VLRRENDMOLEGRAVE, 0) + COALESCE(tot.VLRRENDMOLEGRAVE13, 0) +
            COALESCE(tot.VLRAUXMORADIA, 0) + COALESCE(tot.VLRISENOUTROS, 0) AS ESOCIAL_ISENTOS_TOTAL,
            tot.I_DADOS_EVENTOS
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_TOTALIZADOR tot
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = tot.CODI_EMP AND ext.I_DADOS_EVENTOS = tot.I_DADOS_EVENTOS
        JOIN bethadba.GEEMPRE emp 
          ON emp.CODI_EMP = tot.CODI_EMP
        LEFT JOIN (
            SELECT CODI_EMP, CPF, MAX(NOME) AS NOME
            FROM bethadba.FOEMPREGADOS
            GROUP BY CODI_EMP, CPF
        ) emp_col
          ON emp_col.CODI_EMP = tot.CODI_EMP AND emp_col.CPF = ext.CPF
        JOIN bethadba.FOESOCIAL_DADOS_EVENTOS fde
          ON fde.CODI_EMP = tot.CODI_EMP AND fde.I_DADOS_EVENTOS = tot.I_DADOS_EVENTOS
        WHERE YEAR(tot.COMPETENCIA_REF) = {ano}
          AND fde.VALIDADO = 1
          AND NOT EXISTS (
              SELECT 1 FROM bethadba.FOESOCIAL_DADOS_EVENTOS ret
              WHERE ret.CODI_EMP_EVENTO_RETIFICADO = fde.CODI_EMP
                AND ret.I_DADOS_EVENTOS_EVENTO_RETIFICADO = fde.I_DADOS_EVENTOS
                AND ret.VALIDADO = 1
          )
          {emp_filter_sql}
    """
    cur.execute(sql_es)
    df_es = pd.DataFrame.from_records(cur.fetchall(), columns=[c[0].lower() for c in cur.description])
    print(f"      Total registros eSocial validados retornados: {len(df_es):,}")

    if df_es.empty:
        print(f"Nenhum registro encontrado no eSocial para o ano {ano}.")
        conn.close()
        return None

    # Subtabela de Dependentes do eSocial
    sql_dep = f"""
        SELECT ded.CODI_EMP, ded.COMPETENCIA_REF, ext.CPF,
               SUM(COALESCE(ded.VLRDEDDEP, 0)) AS ESOCIAL_DED_DEPENDENTES
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_DED_DEPEN ded
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = ded.CODI_EMP AND ext.I_DADOS_EVENTOS = ded.I_DADOS_EVENTOS
        WHERE YEAR(ded.COMPETENCIA_REF) = {ano}
        GROUP BY ded.CODI_EMP, ded.COMPETENCIA_REF, ext.CPF
    """
    cur.execute(sql_dep)
    df_dep = pd.DataFrame.from_records(cur.fetchall(), columns=[c[0].lower() for c in cur.description])

    # Subtabela de Plano de Saúde do eSocial
    print("      Extraindo subtabela de Plano de Saúde do eSocial...")
    sql_ps = f"""
        SELECT ps.CODI_EMP, ps.COMPETENCIA_REF, ext.CPF,
               SUM(COALESCE(ps.VLRSAUDETIT, 0)) AS ESOCIAL_PLANO_SAUDE
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_PLANO_SAUDE ps
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = ps.CODI_EMP AND ext.I_DADOS_EVENTOS = ps.I_DADOS_EVENTOS
        WHERE YEAR(ps.COMPETENCIA_REF) = {ano}
        GROUP BY ps.CODI_EMP, ps.COMPETENCIA_REF, ext.CPF
    """
    cur.execute(sql_ps)
    df_ps = pd.DataFrame.from_records(cur.fetchall(), columns=[c[0].lower() for c in cur.description])

    # Agrupar eSocial por Empresa + Competência + CPF
    df_es_grp = df_es.groupby(['codi_emp', 'nome_emp', 'cgce_emp', 'competencia_ref', 'cpf', 'nome_colaborador']).agg({
        'esocial_rendtrib_mensal': 'sum',
        'esocial_rendtrib_13': 'sum',
        'esocial_prev_mensal': 'sum',
        'esocial_prev_13': 'sum',
        'esocial_irrf_mensal': 'sum',
        'esocial_irrf_13': 'sum',
        'esocial_isentos_total': 'sum'
    }).reset_index()

    if not df_dep.empty:
        df_es_grp = pd.merge(df_es_grp, df_dep, on=['codi_emp', 'competencia_ref', 'cpf'], how='left')
    else:
        df_es_grp['esocial_ded_dependentes'] = 0.0
    df_es_grp['esocial_ded_dependentes'] = df_es_grp['esocial_ded_dependentes'].fillna(0.0)

    if not df_ps.empty:
        df_es_grp = pd.merge(df_es_grp, df_ps, on=['codi_emp', 'competencia_ref', 'cpf'], how='left')
    else:
        df_es_grp['esocial_plano_saude'] = 0.0
    df_es_grp['esocial_plano_saude'] = df_es_grp['esocial_plano_saude'].fillna(0.0)

    # 2. Folha do Sistema (Regime de Caixa por Data de Pagamento no ano base)
    print("\n[2/4] Extraindo dados do Sistema (Folha de Pagamento em Regime de Caixa)...")
    empresas = list(df_es['codi_emp'].unique())
    emp_str = ",".join(map(str, empresas))

    # TIPO em FOBASESIRRF: 1=Mensal/Folha Normal, 2=Adiantamento 13o,
    # 3=13o Salario Rescisao, 4=PLR, 6=13o Salario 2a Parcela, 7=RRA
    sql_sis = f"""
        SELECT
            s.CODI_EMP,
            emp.NOME_EMP,
            emp.CGCE_EMP,
            YMD(YEAR(b.DATA_PAGTO), MONTH(b.DATA_PAGTO), 1) AS COMPETENCIA_REF,
            e.CPF,
            MAX(e.NOME) AS NOME_COLABORADOR,
            SUM(CASE WHEN b.TIPO NOT IN (2, 3, 4, 6) THEN b.BASE ELSE 0 END) AS SISTEMA_RENDTRIB_MENSAL,
            SUM(CASE WHEN b.TIPO IN (2, 3, 6) THEN b.BASE ELSE 0 END)         AS SISTEMA_RENDTRIB_13,
            SUM(CASE WHEN b.TIPO = 4 THEN b.BASE ELSE 0 END)               AS SISTEMA_RENDTRIB_PLR,
            SUM(CASE WHEN b.TIPO NOT IN (2, 3, 4, 6) THEN b.ABATIMENTOS ELSE 0 END) AS SISTEMA_PREV_MENSAL,
            SUM(CASE WHEN b.TIPO IN (2, 3, 6) THEN b.ABATIMENTOS ELSE 0 END)         AS SISTEMA_PREV_13,
            SUM(b.DEPEND_DESCONTO)                                                AS SISTEMA_DED_DEPENDENTES,
            SUM(CASE WHEN b.TIPO NOT IN (2, 3, 4, 6) THEN b.VALOR ELSE 0 END) AS SISTEMA_IRRF_MENSAL,
            SUM(CASE WHEN b.TIPO IN (2, 3, 6) THEN b.VALOR ELSE 0 END)         AS SISTEMA_IRRF_13,
            SUM(CASE WHEN b.TIPO = 4 THEN b.VALOR ELSE 0 END)               AS SISTEMA_IRRF_PLR
        FROM bethadba.FOBASESSERVIRRF s
        JOIN bethadba.FOBASESIRRF b ON b.I_BASESIRRF = s.I_BASESIRRF
        JOIN bethadba.FOEMPREGADOS e ON e.CODI_EMP = s.CODI_EMP AND e.I_EMPREGADOS = s.I_EMPREGADOS
        JOIN bethadba.GEEMPRE emp ON emp.CODI_EMP = s.CODI_EMP
        WHERE s.CODI_EMP IN ({emp_str})
          AND s.RATEIO = 0
          AND b.DATA_PAGTO >= '{ano}-01-01'
          AND b.DATA_PAGTO <= '{ano}-12-31'
        GROUP BY
            s.CODI_EMP, emp.NOME_EMP, emp.CGCE_EMP,
            YMD(YEAR(b.DATA_PAGTO), MONTH(b.DATA_PAGTO), 1),
            e.CPF
    """
    cur.execute(sql_sis)
    df_sis = pd.DataFrame.from_records(cur.fetchall(), columns=[c[0].lower() for c in cur.description])
    print(f"      Total registros Sistema retornados: {len(df_sis):,}")

    # Eventos complementares do Sistema: Plano de Saúde (CLASSIFICACAO=29) e Isentos
    print("      Extraindo eventos de Plano de Saúde e Isentos do Sistema...")
    sql_eve = f"""
        SELECT
            bs.CODI_EMP,
            YMD(YEAR(bs.DATA_PAGTO), MONTH(bs.DATA_PAGTO), 1) AS COMPETENCIA_REF,
            e.CPF,
            SUM(CASE WHEN ev.CLASSIFICACAO = 29 THEN
                    CASE WHEN m.PROV_DESC IN ('D', 'ID') THEN m.VALOR_CAL ELSE -m.VALOR_CAL END
                ELSE 0 END) AS SISTEMA_PLANO_SAUDE,
            SUM(CASE WHEN ev.REND_ISENTOS > 0 OR ev.CLASSIFICACAO = 46
                     OR ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (72, 73) THEN
                    CASE WHEN m.PROV_DESC IN ('P', 'I') THEN m.VALOR_CAL ELSE -m.VALOR_CAL END
                ELSE 0 END) AS SISTEMA_ISENTOS_TOTAL
        FROM bethadba.FOBASESSERV bs
        JOIN bethadba.FOMOVTOSERV m
          ON m.CODI_EMP = bs.CODI_EMP AND m.I_CALCULOS = bs.I_CALCULOS
        JOIN bethadba.FOEMPREGADOS e
          ON e.CODI_EMP = bs.CODI_EMP AND e.I_EMPREGADOS = bs.I_EMPREGADOS
        JOIN bethadba.FOPARMTO p
          ON p.CODI_EMP = bs.CODI_EMP
        JOIN bethadba.FOEVENTOS ev
          ON ev.CODI_EMP = p.CODI_EMP_EVE AND ev.I_EVENTOS = m.I_EVENTOS
        WHERE bs.CODI_EMP IN ({emp_str})
          AND m.RATEIO = 0
          AND m.ORIGEM <> 'F'
          AND ev.SOMA_INF_REN = 'S'
          AND bs.DATA_PAGTO >= '{ano}-01-01'
          AND bs.DATA_PAGTO <= '{ano}-12-31'
        GROUP BY
            bs.CODI_EMP,
            YMD(YEAR(bs.DATA_PAGTO), MONTH(bs.DATA_PAGTO), 1),
            e.CPF
    """
    cur.execute(sql_eve)
    df_eve = pd.DataFrame.from_records(cur.fetchall(), columns=[c[0].lower() for c in cur.description])
    print(f"      Total registros de Plano de Saúde/Isentos do Sistema: {len(df_eve):,}")
    conn.close()

    if not df_eve.empty:
        df_sis = pd.merge(df_sis, df_eve, on=['codi_emp', 'competencia_ref', 'cpf'], how='left')
    else:
        df_sis['sistema_plano_saude'] = 0.0
        df_sis['sistema_isentos_total'] = 0.0
    df_sis['sistema_plano_saude'] = df_sis['sistema_plano_saude'].fillna(0.0)
    df_sis['sistema_isentos_total'] = df_sis['sistema_isentos_total'].fillna(0.0)

    # 3. Conciliação e Cruzamento Exato
    print("\n[3/4] Realizando Cruzamento e Classificação das Divergências...")
    m = pd.merge(df_es_grp, df_sis, on=['codi_emp', 'competencia_ref', 'cpf'], how='outer')

    # Consolidar dados cadastrais
    m['nome_emp'] = m['nome_emp_x'].combine_first(m['nome_emp_y'])
    m['cgce_emp'] = m['cgce_emp_x'].combine_first(m['cgce_emp_y'])
    m['nome_colaborador'] = m['nome_colaborador_x'].combine_first(m['nome_colaborador_y'])
    m.drop(columns=['nome_emp_x', 'nome_emp_y', 'cgce_emp_x', 'cgce_emp_y', 'nome_colaborador_x', 'nome_colaborador_y'], inplace=True)

    # Converter numéricos
    cols_num = [c for c in m.columns if c.startswith(('esocial_', 'sistema_'))]
    for c in cols_num:
        m[c] = pd.to_numeric(m[c], errors='coerce').fillna(0.0).astype(float).round(2)

    # Totais consolidados
    m['esocial_rendtrib_total'] = (m['esocial_rendtrib_mensal'] + m['esocial_rendtrib_13']).round(2)
    m['sistema_rendtrib_total'] = (m['sistema_rendtrib_mensal'] + m['sistema_rendtrib_13'] + m['sistema_rendtrib_plr']).round(2)
    m['dif_rendtrib'] = (m['sistema_rendtrib_total'] - m['esocial_rendtrib_total']).round(2)

    m['esocial_prev_total'] = (m['esocial_prev_mensal'] + m['esocial_prev_13']).round(2)
    m['sistema_prev_total'] = (m['sistema_prev_mensal'] + m['sistema_prev_13']).round(2)
    m['dif_previdencia'] = (m['sistema_prev_total'] - m['esocial_prev_total']).round(2)

    m['esocial_irrf_total'] = (m['esocial_irrf_mensal'] + m['esocial_irrf_13']).round(2)
    m['sistema_irrf_total'] = (m['sistema_irrf_mensal'] + m['sistema_irrf_13'] + m['sistema_irrf_plr']).round(2)
    m['dif_irrf'] = (m['sistema_irrf_total'] - m['esocial_irrf_total']).round(2)

    m['dif_dependentes'] = (m['sistema_ded_dependentes'] - m['esocial_ded_dependentes']).round(2)
    m['dif_plano_saude'] = (m['sistema_plano_saude'] - m['esocial_plano_saude']).round(2)
    m['dif_isentos'] = (m['sistema_isentos_total'] - m['esocial_isentos_total']).round(2)

    # Regra de Classificação de Status
    def classificar(r):
        tem_esoc = (r['esocial_rendtrib_total'] > 0 or r['esocial_prev_total'] > 0 or r['esocial_irrf_total'] > 0
                    or r['esocial_plano_saude'] != 0 or r['esocial_isentos_total'] != 0)
        tem_sis = (r['sistema_rendtrib_total'] > 0 or r['sistema_prev_total'] > 0 or r['sistema_irrf_total'] > 0
                   or r['sistema_plano_saude'] != 0 or r['sistema_isentos_total'] != 0)

        if tem_esoc and not tem_sis:
            return "APENAS NO eSOCIAL", "Registro presente no eSocial, porém sem cálculo de folha correspondente no Domínio."
        if tem_sis and not tem_esoc:
            return "PENDÊNCIA eSOCIAL", "Folha calculada no Domínio, mas sem retorno do evento S-5002 no Extrator da DIRF (evento pendente de envio ou retorno)."

        dif_irrf = abs(r['dif_irrf']) > 0.01
        dif_rend = abs(r['dif_rendtrib']) > 0.01
        dif_prev = abs(r['dif_previdencia']) > 0.01
        dif_dep = abs(r['dif_dependentes']) > 0.01
        dif_ps = abs(r['dif_plano_saude']) > 0.01
        dif_isen = abs(r['dif_isentos']) > 0.01

        if dif_irrf:
            return "DIVERGÊNCIA CRÍTICA: IRRF", f"Divergência de IRRF Retido: Folha R$ {r['sistema_irrf_total']:,.2f} x eSocial R$ {r['esocial_irrf_total']:,.2f} (Diferença: R$ {r['dif_irrf']:,.2f})"
        if dif_rend:
            return "DIVERGÊNCIA: BASE DE CÁLCULO", f"Divergência na Base Tributável: Folha R$ {r['sistema_rendtrib_total']:,.2f} x eSocial R$ {r['esocial_rendtrib_total']:,.2f} (Diferença: R$ {r['dif_rendtrib']:,.2f})"
        if dif_prev or dif_dep:
            motivos = []
            if dif_prev:
                motivos.append(f"INSS: Folha R$ {r['sistema_prev_total']:,.2f} x eSoc R$ {r['esocial_prev_total']:,.2f}")
            if dif_dep:
                motivos.append(f"Dependentes: Folha R$ {r['sistema_ded_dependentes']:,.2f} x eSoc R$ {r['esocial_ded_dependentes']:,.2f}")
            return "DIVERGÊNCIA: DEDUÇÕES (INSS/DEP)", " | ".join(motivos)
        if dif_ps or dif_isen:
            motivos = []
            if dif_ps:
                motivos.append(f"Plano Saúde: Folha R$ {r['sistema_plano_saude']:,.2f} x eSoc R$ {r['esocial_plano_saude']:,.2f}")
            if dif_isen:
                motivos.append(f"Isentos: Folha R$ {r['sistema_isentos_total']:,.2f} x eSoc R$ {r['esocial_isentos_total']:,.2f}")
            return "DIVERGÊNCIA: PLANO SAÚDE/ISENTOS", " | ".join(motivos)

        return "CONCILIADO 100%", "Valores idênticos centavo a centavo entre Domínio e eSocial."

    res_class = m.apply(classificar, axis=1)
    m['status_categoria'] = [x[0] for x in res_class]
    m['detalhes_explicativos'] = [x[1] for x in res_class]

    # Formatar CPFs e CNPJs para exibição
    m['cpf_formatado'] = m['cpf'].apply(format_cpf)
    m['cnpj_formatado'] = m['cgce_emp'].apply(format_cnpj)

    return m


def aplicar_estilos_gerais(ws, title, is_dashboard=False):
    # Cores
    cor_azul_navy = "1B365D"
    cor_branco = "FFFFFF"
    cor_cinza_claro = "F7FAFC"
    cor_cinza_borda = "E2E8F0"
    
    font_header = Font(name="Segoe UI", size=10, bold=True, color=cor_branco)
    fill_header = PatternFill(start_color=cor_azul_navy, end_color=cor_azul_navy, fill_type="solid")
    border_thin = Side(border_style="thin", color=cor_cinza_borda)
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    if not is_dashboard:
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 28

        for cell in ws[1]:
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border

        headers_upper = [str(c.value or "").upper() for c in ws[1]]

        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 20
            is_even = (row_idx % 2 == 0)

            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = cell_border
                header_name = headers_upper[col_idx - 1]
                val = cell.value

                # Zebra
                if is_even:
                    cell.fill = PatternFill(start_color=cor_cinza_claro, end_color=cor_cinza_claro, fill_type="solid")

                # Formatações de valores
                if isinstance(val, (int, float, Decimal)):
                    if any(k in header_name for k in ["VALOR", "REND", "PREV", "IRRF", "DIF", "TOTAL", "ISENT", "SAUDE", "DEP", "BASE"]):
                        cell.number_format = '"R$ "#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        if "DIF" in header_name and abs(float(val)) > 0.01:
                            if "IRRF" in header_name:
                                cell.font = Font(name="Segoe UI", size=10, bold=True, color="9C1C1C")
                                cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                            else:
                                cell.font = Font(name="Segoe UI", size=10, bold=True, color="B45309")
                                cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    elif any(k in header_name for k in ["QTD", "COLABORADORES", "TOTAL"]):
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                elif isinstance(val, (datetime, date)):
                    cell.number_format = 'mm/yyyy'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    str_v = str(val or "").strip()
                    if "CRÍTICA" in str_v or str_v == "DIVERGÊNCIA CRÍTICA: IRRF":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="9C1C1C")
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif "DIVERGÊNCIA" in str_v:
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="B45309")
                        cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif "PENDÊNCIA" in str_v:
                        cell.font = Font(name="Segoe UI", size=10, color="4B5563")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif "CONCILIADO" in str_v or str_v == "OK":
                        cell.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
                        cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif any(k in header_name for k in ["CPF", "CNPJ", "COMPET", "DATA"]):
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-ajuste de colunas
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val = cell.value
                if val is not None:
                    if isinstance(val, (int, float, Decimal)):
                        s = f"R$ {val:,.2f}"
                    elif isinstance(val, (datetime, date)):
                        s = "00/2026"
                    else:
                        s = str(val)
                    max_len = max(max_len, len(s))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        ws.auto_filter.ref = ws.dimensions


def gerar_relatorio_executivo_excel(df, output_path="relatorios/Auditoria_Extrator_DIRF_eSocial_2026.xlsx"):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    print(f"\n[4/4] Gerando Planilha Executiva em {output_path}...")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove aba padrao

    # =========================================================================
    # 1. ABA DASHBOARD & RESUMO EXECUTIVO
    # =========================================================================
    ws_dash = wb.create_sheet(title="📊 Resumo Executivo")
    ws_dash.views.sheetView[0].showGridLines = True

    # Métricas Gerais
    total_linhas = len(df)
    total_empresas = df['codi_emp'].nunique()
    total_cpfs = df['cpf'].nunique()

    conciliados = (df['status_categoria'] == "CONCILIADO 100%").sum()
    div_irrf = (df['status_categoria'] == "DIVERGÊNCIA CRÍTICA: IRRF").sum()
    div_base = (df['status_categoria'] == "DIVERGÊNCIA: BASE DE CÁLCULO").sum()
    div_ded = (df['status_categoria'] == "DIVERGÊNCIA: DEDUÇÕES (INSS/DEP)").sum()
    div_ps = (df['status_categoria'] == "DIVERGÊNCIA: PLANO SAÚDE/ISENTOS").sum()
    pend_esoc = (df['status_categoria'] == "PENDÊNCIA eSOCIAL").sum()
    apenas_esoc = (df['status_categoria'] == "APENAS NO eSOCIAL").sum()

    # Os totais financeiros abaixo consideram APENAS colaboradores com folha
    # calculada E retorno eSocial já processado (exclui PENDÊNCIA/APENAS NO
    # eSOCIAL), para não misturar atraso operacional de processamento com
    # divergência fiscal real na leitura do dashboard.
    df_reconciliavel = df[~df['status_categoria'].isin(["PENDÊNCIA eSOCIAL", "APENAS NO eSOCIAL"])]

    tot_sis_rend = df_reconciliavel['sistema_rendtrib_total'].sum()
    tot_esoc_rend = df_reconciliavel['esocial_rendtrib_total'].sum()
    tot_dif_rend = df_reconciliavel['dif_rendtrib'].sum()

    tot_sis_irrf = df_reconciliavel['sistema_irrf_total'].sum()
    tot_esoc_irrf = df_reconciliavel['esocial_irrf_total'].sum()
    tot_dif_irrf = df_reconciliavel['dif_irrf'].sum()

    pend_sis_rend = df.loc[df['status_categoria'] == "PENDÊNCIA eSOCIAL", 'sistema_rendtrib_total'].sum()

    # Título do Dashboard
    ws_dash.merge_cells("B2:J2")
    ws_dash["B2"] = "AUDITORIA DE FIDEDIGNIDADE: EXTRATOR DA DIRF x eSOCIAL (ANO BASE 2026)"
    ws_dash["B2"].font = Font(name="Segoe UI", size=16, bold=True, color="1B365D")
    ws_dash["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[2].height = 40

    ws_dash.merge_cells("B3:J3")
    ws_dash["B3"] = f"Relatório Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Base de Dados: Domínio Sistemas Oficial | Universo Auditado: {total_empresas} Empresas / {total_cpfs:,} Colaboradores"
    ws_dash["B3"].font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
    ws_dash["B3"].alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[3].height = 20

    # Cards de KPIs (Linha 5-6: visao geral | Linha 7-8: detalhamento de divergencias)
    kpis = [
        ("B5", "C6", "EMPRESAS AUDITADAS", f"{total_empresas}", "1B365D", "F0F4F8"),
        ("D5", "E6", "100% CONCILIADOS", f"{conciliados:,} ({conciliados/total_linhas*100:.1f}%)", "065F46", "D1FAE5"),
        ("F5", "G6", "DIVERGÊNCIA CRÍTICA IRRF", f"{div_irrf}", "9C1C1C", "FEE2E2"),
        ("H5", "I6", "DIVERGÊNCIA DE BASE", f"{div_base:,}", "B45309", "FEF3C7"),
        ("J5", "J6", "PENDÊNCIAS eSOCIAL", f"{pend_esoc:,}", "374151", "F3F4F6"),
        ("B7", "C8", "DIVERGÊNCIA DEDUÇÕES (INSS/DEP)", f"{div_ded:,}", "B45309", "FEF3C7"),
        ("D7", "E8", "DIVERGÊNCIA PLANO SAÚDE/ISENTOS", f"{div_ps:,}", "B45309", "FEF3C7"),
        ("F7", "G8", "APENAS NO eSOCIAL", f"{apenas_esoc:,}", "374151", "F3F4F6"),
    ]

    for top_l, bot_r, label, val_str, text_color, bg_color in kpis:
        top_col = top_l[0]
        bot_col = bot_r[0]
        top_row = int(top_l[1:])
        bot_row = int(bot_r[1:])

        ws_dash.merge_cells(f"{top_l}:{bot_r}")
        cell = ws_dash[top_l]
        cell.value = f"{label}\n{val_str}"
        cell.font = Font(name="Segoe UI", size=12, bold=True, color=text_color)
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin = Side(border_style="medium", color=text_color)
        box = Border(left=thin, right=thin, top=thin, bottom=thin)
        for r in range(top_row, bot_row + 1):
            for c in range(ord(top_col) - ord('A') + 1, ord(bot_col) - ord('A') + 2):
                ws_dash.cell(row=r, column=c).border = box

    for row_h in (5, 6, 7, 8):
        ws_dash.row_dimensions[row_h].height = 26

    # Tabela 1: Resumo dos Totais Financeiros
    ws_dash["B10"] = "1. CONCILIAÇÃO FINANCEIRA GLOBAL (ANO 2026)"
    ws_dash["B10"].font = Font(name="Segoe UI", size=11, bold=True, color="1B365D")

    headers_fin = ["RUBRICA FISCAL", "TOTAL SISTEMA (FOLHA)", "TOTAL eSOCIAL (S-5002)", "DIFERENÇA LÍQUIDA", "STATUS"]
    for c_idx, h in enumerate(headers_fin, start=2):
        cell = ws_dash.cell(row=11, column=c_idx, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    linhas_fin = [
        ("Rendimentos Tributáveis", tot_sis_rend, tot_esoc_rend, tot_dif_rend),
        ("Previdência Oficial (INSS)", df_reconciliavel['sistema_prev_total'].sum(), df_reconciliavel['esocial_prev_total'].sum(), df_reconciliavel['dif_previdencia'].sum()),
        ("Dedução Dependentes", df_reconciliavel['sistema_ded_dependentes'].sum(), df_reconciliavel['esocial_ded_dependentes'].sum(), df_reconciliavel['dif_dependentes'].sum()),
        ("Plano de Saúde", df_reconciliavel['sistema_plano_saude'].sum(), df_reconciliavel['esocial_plano_saude'].sum(), df_reconciliavel['dif_plano_saude'].sum()),
        ("Rendimentos Isentos/Não Tributáveis", df_reconciliavel['sistema_isentos_total'].sum(), df_reconciliavel['esocial_isentos_total'].sum(), df_reconciliavel['dif_isentos'].sum()),
        ("IRRF Retido na Fonte", tot_sis_irrf, tot_esoc_irrf, tot_dif_irrf),
    ]

    for r_idx, (rub, v_sis, v_esoc, v_dif) in enumerate(linhas_fin, start=12):
        ws_dash.cell(row=r_idx, column=2, value=rub).font = Font(name="Segoe UI", size=10, bold=True)
        c_sis = ws_dash.cell(row=r_idx, column=3, value=v_sis)
        c_es = ws_dash.cell(row=r_idx, column=4, value=v_esoc)
        c_dif = ws_dash.cell(row=r_idx, column=5, value=v_dif)

        for c in [c_sis, c_es, c_dif]:
            c.number_format = '"R$ "#,##0.00'
            c.alignment = Alignment(horizontal="right", vertical="center")

        c_st = ws_dash.cell(row=r_idx, column=6)
        if abs(v_dif) < 1.0:
            c_st.value = "CONCILIADO"
            c_st.font = Font(name="Segoe UI", size=10, bold=True, color="065F46")
            c_st.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        else:
            c_st.value = "COM DIVERGÊNCIA"
            c_st.font = Font(name="Segoe UI", size=10, bold=True, color="9C1C1C")
            c_st.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        c_st.alignment = Alignment(horizontal="center", vertical="center")

    nota_row = 12 + len(linhas_fin)
    ws_dash.merge_cells(f"B{nota_row}:J{nota_row}")
    ws_dash[f"B{nota_row}"] = (
        f"Nota: este quadro exclui {pend_esoc:,} colaboradores em PENDÊNCIA eSOCIAL "
        f"(R$ {pend_sis_rend:,.2f} em folha ainda sem retorno S-5002) e {apenas_esoc:,} registros "
        f"APENAS NO eSOCIAL, para não misturar atraso de processamento com divergência fiscal real. "
        f"Ver abas '⏳ Pendências eSocial' e '🏳️ Apenas no eSocial'."
    )
    ws_dash[f"B{nota_row}"].font = Font(name="Segoe UI", size=9, italic=True, color="4B5563")
    ws_dash[f"B{nota_row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_dash.row_dimensions[nota_row].height = 28

    # Tabela 2: Guia de Ação para o Usuário
    guia_header_row = nota_row + 2
    ws_dash[f"B{guia_header_row}"] = "2. GUIA DE ORIENTAÇÃO E AÇÕES RECOMENDADAS"
    ws_dash[f"B{guia_header_row}"].font = Font(name="Segoe UI", size=11, bold=True, color="1B365D")

    headers_guia = ["CATEGORIA", "IMPACTO / RISCO", "QTD CASOS", "AÇÃO OPERACIONAL RECOMENDADA"]
    for c_idx, h in enumerate(headers_guia, start=2):
        cell = ws_dash.cell(row=guia_header_row + 1, column=c_idx, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    acoes = [
        ("DIVERGÊNCIA CRÍTICA: IRRF", "ALTÍSSIMO (Risco de Autuação Fiscal e Malha da DIRF)", div_irrf,
         "Acessar a aba '🚨 Divergências Críticas IRRF'. Verificar retenções de sócios/diretores não enviadas no S-1210 e retificar o evento."),
        ("DIVERGÊNCIA: BASE DE CÁLCULO", "ALTO (Malha fina no IRPF do colaborador)", div_base,
         "Acessar a aba '⚖️ Divergências de Base'. Conferir duplicidades de rubricas, fechamento de folha ou cálculo de rescisão sem envio do S-1200."),
        ("DIVERGÊNCIA: DEDUÇÕES", "MÉDIO (Diferença de INSS ou Dependentes)", div_ded,
         "Verificar cadastro de dependentes para fins de IRRF e tabela de incidência da previdência nas rubricas."),
        ("DIVERGÊNCIA: PLANO SAÚDE/ISENTOS", "MÉDIO (Diferença em Plano de Saúde ou Rendimentos Isentos)", div_ps,
         "Acessar a aba '🩺 Divergências Plano Saúde'. Conferir eventos com CLASSIFICACAO=29 (plano de saúde) e rubricas isentas (rescisão, moléstia grave, diárias) não refletidas no S-5002."),
        ("PENDÊNCIA eSOCIAL", "OPERACIONAL (S-5002 não retornado)", pend_esoc,
         "Acessar '⏳ Pendências eSocial'. Consultar na Domínio se o evento S-1210 foi transmitido e reprocessar o Extrator da DIRF para importar os totalizadores."),
        ("APENAS NO eSOCIAL", "OPERACIONAL (evento eSocial sem folha correspondente)", apenas_esoc,
         "Acessar '🏳️ Apenas no eSocial'. Verificar se a folha foi calculada em outra competência ou se o evento foi enviado indevidamente."),
        ("CONCILIADO 100%", "ZERO (Auditoria perfeita centavo a centavo)", conciliados,
         "Nenhuma ação necessária. Dados 100% fidedignos e validados."),
    ]

    for r_idx, (cat, risco, qtd, acao) in enumerate(acoes, start=guia_header_row + 2):
        c1 = ws_dash.cell(row=r_idx, column=2, value=cat)
        c2 = ws_dash.cell(row=r_idx, column=3, value=risco)
        c3 = ws_dash.cell(row=r_idx, column=4, value=qtd)
        c4 = ws_dash.cell(row=r_idx, column=5, value=acao)

        c1.font = Font(name="Segoe UI", size=9, bold=True)
        c2.font = Font(name="Segoe UI", size=9)
        c3.font = Font(name="Segoe UI", size=9, bold=True)
        c3.number_format = '#,##0'
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c4.font = Font(name="Segoe UI", size=9)

    # Ajustar larguras das colunas do Dashboard
    ws_dash.column_dimensions["A"].width = 4
    ws_dash.column_dimensions["B"].width = 32
    ws_dash.column_dimensions["C"].width = 24
    ws_dash.column_dimensions["D"].width = 24
    ws_dash.column_dimensions["E"].width = 50
    ws_dash.column_dimensions["F"].width = 24
    ws_dash.column_dimensions["G"].width = 24
    ws_dash.column_dimensions["H"].width = 24
    ws_dash.column_dimensions["I"].width = 24
    ws_dash.column_dimensions["J"].width = 24

    # =========================================================================
    # 2. ABA DIVERGÊNCIAS CRÍTICAS DE IRRF (PRIORIDADE 1)
    # =========================================================================
    df_irrf = df[df['status_categoria'] == "DIVERGÊNCIA CRÍTICA: IRRF"].copy()
    cols_irrf = [
        'codi_emp', 'nome_emp', 'cnpj_formatado', 'competencia_ref', 'cpf_formatado', 'nome_colaborador',
        'sistema_irrf_total', 'esocial_irrf_total', 'dif_irrf',
        'sistema_rendtrib_total', 'esocial_rendtrib_total', 'dif_rendtrib',
        'detalhes_explicativos'
    ]
    df_irrf_exp = df_irrf[cols_irrf].copy()
    df_irrf_exp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'COMPETÊNCIA', 'CPF', 'COLABORADOR',
        'IRRF SISTEMA (FOLHA)', 'IRRF eSOCIAL (S-5002)', 'DIFERENÇA IRRF',
        'REND. TRIB. SISTEMA', 'REND. TRIB. eSOCIAL', 'DIFERENÇA RENDIMENTOS',
        'DIAGNÓSTICO DA DIVERGÊNCIA'
    ]

    ws_irrf = wb.create_sheet(title="🚨 Divergências Críticas IRRF")
    ws_irrf.append(list(df_irrf_exp.columns))
    for r in df_irrf_exp.itertuples(index=False):
        ws_irrf.append(list(r))
    aplicar_estilos_gerais(ws_irrf, "Divergencias Criticas IRRF")

    # =========================================================================
    # 3. ABA DIVERGÊNCIAS DE BASE DE CÁLCULO (PRIORIDADE 2)
    # =========================================================================
    df_base_rows = df[df['status_categoria'] == "DIVERGÊNCIA: BASE DE CÁLCULO"].copy()
    cols_base = [
        'codi_emp', 'nome_emp', 'cnpj_formatado', 'competencia_ref', 'cpf_formatado', 'nome_colaborador',
        'sistema_rendtrib_mensal', 'esocial_rendtrib_mensal',
        'sistema_rendtrib_13', 'esocial_rendtrib_13',
        'sistema_rendtrib_total', 'esocial_rendtrib_total', 'dif_rendtrib',
        'sistema_irrf_total', 'esocial_irrf_total',
        'detalhes_explicativos'
    ]
    df_base_exp = df_base_rows[cols_base].copy()
    df_base_exp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'COMPETÊNCIA', 'CPF', 'COLABORADOR',
        'FOLHA MENSAL (SISTEMA)', 'FOLHA MENSAL (eSOCIAL)',
        '13º SALÁRIO (SISTEMA)', '13º SALÁRIO (eSOCIAL)',
        'TOTAL REND. TRIB. SISTEMA', 'TOTAL REND. TRIB. eSOCIAL', 'DIFERENÇA RENDIMENTOS',
        'IRRF SISTEMA', 'IRRF eSOCIAL',
        'DIAGNÓSTICO DA DIVERGÊNCIA'
    ]

    ws_base = wb.create_sheet(title="⚖️ Divergências de Base")
    ws_base.append(list(df_base_exp.columns))
    for r in df_base_exp.itertuples(index=False):
        ws_base.append(list(r))
    aplicar_estilos_gerais(ws_base, "Divergencias de Base")

    # =========================================================================
    # 4. ABA PENDÊNCIAS eSOCIAL (PRIORIDADE 3)
    # =========================================================================
    df_pend = df[df['status_categoria'] == "PENDÊNCIA eSOCIAL"].copy()
    cols_pend = [
        'codi_emp', 'nome_emp', 'cnpj_formatado', 'competencia_ref', 'cpf_formatado', 'nome_colaborador',
        'sistema_rendtrib_total', 'sistema_prev_total', 'sistema_irrf_total',
        'detalhes_explicativos'
    ]
    df_pend_exp = df_pend[cols_pend].copy()
    df_pend_exp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'COMPETÊNCIA', 'CPF', 'COLABORADOR',
        'REND. TRIB. SISTEMA', 'INSS SISTEMA', 'IRRF SISTEMA',
        'SITUAÇÃO / MOTIVO'
    ]

    ws_pend = wb.create_sheet(title="⏳ Pendências eSocial")
    ws_pend.append(list(df_pend_exp.columns))
    for r in df_pend_exp.itertuples(index=False):
        ws_pend.append(list(r))
    aplicar_estilos_gerais(ws_pend, "Pendencias eSocial")

    # =========================================================================
    # 4b. ABA APENAS NO eSOCIAL
    # =========================================================================
    df_apenas = df[df['status_categoria'] == "APENAS NO eSOCIAL"].copy()
    cols_apenas = [
        'codi_emp', 'nome_emp', 'cnpj_formatado', 'competencia_ref', 'cpf_formatado', 'nome_colaborador',
        'esocial_rendtrib_total', 'esocial_prev_total', 'esocial_irrf_total',
        'detalhes_explicativos'
    ]
    df_apenas_exp = df_apenas[cols_apenas].copy()
    df_apenas_exp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'COMPETÊNCIA', 'CPF', 'COLABORADOR',
        'REND. TRIB. eSOCIAL', 'INSS eSOCIAL', 'IRRF eSOCIAL',
        'SITUAÇÃO / MOTIVO'
    ]

    ws_apenas = wb.create_sheet(title="🏳️ Apenas no eSocial")
    ws_apenas.append(list(df_apenas_exp.columns))
    for r in df_apenas_exp.itertuples(index=False):
        ws_apenas.append(list(r))
    aplicar_estilos_gerais(ws_apenas, "Apenas no eSocial")

    # =========================================================================
    # 4c. ABA DIVERGÊNCIAS DE PLANO DE SAÚDE / ISENTOS
    # =========================================================================
    df_ps_rows = df[df['status_categoria'] == "DIVERGÊNCIA: PLANO SAÚDE/ISENTOS"].copy()
    cols_ps = [
        'codi_emp', 'nome_emp', 'cnpj_formatado', 'competencia_ref', 'cpf_formatado', 'nome_colaborador',
        'sistema_plano_saude', 'esocial_plano_saude', 'dif_plano_saude',
        'sistema_isentos_total', 'esocial_isentos_total', 'dif_isentos',
        'detalhes_explicativos'
    ]
    df_ps_exp = df_ps_rows[cols_ps].copy()
    df_ps_exp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'COMPETÊNCIA', 'CPF', 'COLABORADOR',
        'PLANO SAÚDE (SISTEMA)', 'PLANO SAÚDE (eSOCIAL)', 'DIFERENÇA PLANO SAÚDE',
        'ISENTOS (SISTEMA)', 'ISENTOS (eSOCIAL)', 'DIFERENÇA ISENTOS',
        'DIAGNÓSTICO DA DIVERGÊNCIA'
    ]

    ws_ps = wb.create_sheet(title="🩺 Divergências Plano Saúde")
    ws_ps.append(list(df_ps_exp.columns))
    for r in df_ps_exp.itertuples(index=False):
        ws_ps.append(list(r))
    aplicar_estilos_gerais(ws_ps, "Divergencias Plano Saude Isentos")

    # =========================================================================
    # 5. ABA RESUMO SINTÉTICO POR EMPRESA
    # =========================================================================
    resumo_emp = df.groupby(['codi_emp', 'nome_emp', 'cnpj_formatado']).agg({
        'cpf': 'nunique',
        'status_categoria': [
            lambda s: (s == "CONCILIADO 100%").sum(),
            lambda s: (s == "DIVERGÊNCIA CRÍTICA: IRRF").sum(),
            lambda s: (s == "DIVERGÊNCIA: BASE DE CÁLCULO").sum(),
            lambda s: (s == "DIVERGÊNCIA: DEDUÇÕES (INSS/DEP)").sum(),
            lambda s: (s == "DIVERGÊNCIA: PLANO SAÚDE/ISENTOS").sum(),
            lambda s: (s == "PENDÊNCIA eSOCIAL").sum(),
        ],
        'sistema_rendtrib_total': 'sum',
        'esocial_rendtrib_total': 'sum',
        'dif_rendtrib': 'sum',
        'sistema_irrf_total': 'sum',
        'esocial_irrf_total': 'sum',
        'dif_irrf': 'sum'
    }).reset_index()

    resumo_emp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'TOTAL COLABORADORES',
        'QTD CONCILIADOS', 'DIVERGÊNCIAS IRRF', 'DIVERGÊNCIAS BASE', 'DIVERGÊNCIAS DEDUÇÕES',
        'DIVERGÊNCIAS PLANO SAÚDE/ISENTOS', 'PENDÊNCIAS eSOCIAL',
        'REND. TRIB. SISTEMA', 'REND. TRIB. eSOCIAL', 'DIF. RENDIMENTOS',
        'IRRF SISTEMA', 'IRRF eSOCIAL', 'DIF. IRRF'
    ]

    def situacao_empresa(r):
        if r['DIVERGÊNCIAS IRRF'] > 0:
            return "CRÍTICO: DIVERGÊNCIA IRRF"
        if r['DIVERGÊNCIAS BASE'] > 0:
            return "ALERTA: DIVERGÊNCIA BASE"
        if r['DIVERGÊNCIAS DEDUÇÕES'] > 0 or r['DIVERGÊNCIAS PLANO SAÚDE/ISENTOS'] > 0:
            return "ALERTA: DIVERGÊNCIA DEDUÇÕES/ISENTOS"
        if r['PENDÊNCIAS eSOCIAL'] > 0:
            return "PENDENTE: eSOCIAL NÃO RETORNADO"
        return "100% CONCILIADO"

    resumo_emp['STATUS GERAL'] = resumo_emp.apply(situacao_empresa, axis=1)

    # Ordenar por criticidade
    ordem_status = {
        "CRÍTICO: DIVERGÊNCIA IRRF": 1,
        "ALERTA: DIVERGÊNCIA BASE": 2,
        "ALERTA: DIVERGÊNCIA DEDUÇÕES/ISENTOS": 3,
        "PENDENTE: eSOCIAL NÃO RETORNADO": 4,
        "100% CONCILIADO": 5
    }
    resumo_emp['ordem'] = resumo_emp['STATUS GERAL'].map(ordem_status)
    resumo_emp.sort_values(by=['ordem', 'DIF. IRRF', 'CÓD. EMP'], ascending=[True, False, True], inplace=True)
    resumo_emp.drop(columns=['ordem'], inplace=True)

    # Reposicionar STATUS GERAL para o início
    cols_order = ['CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'STATUS GERAL', 'TOTAL COLABORADORES',
                  'DIVERGÊNCIAS IRRF', 'DIVERGÊNCIAS BASE', 'DIVERGÊNCIAS DEDUÇÕES',
                  'DIVERGÊNCIAS PLANO SAÚDE/ISENTOS', 'PENDÊNCIAS eSOCIAL', 'QTD CONCILIADOS',
                  'REND. TRIB. SISTEMA', 'REND. TRIB. eSOCIAL', 'DIF. RENDIMENTOS',
                  'IRRF SISTEMA', 'IRRF eSOCIAL', 'DIF. IRRF']
    resumo_emp = resumo_emp[cols_order]

    ws_emp = wb.create_sheet(title="🏢 Resumo por Empresa")
    ws_emp.append(list(resumo_emp.columns))
    for r in resumo_emp.itertuples(index=False):
        ws_emp.append(list(r))
    aplicar_estilos_gerais(ws_emp, "Resumo por Empresa")

    # =========================================================================
    # 6. ABA CONCILIAÇÃO ANALÍTICA COMPLETA
    # =========================================================================
    cols_all = [
        'codi_emp', 'nome_emp', 'cnpj_formatado', 'competencia_ref', 'cpf_formatado', 'nome_colaborador',
        'status_categoria', 'detalhes_explicativos',
        'sistema_rendtrib_total', 'esocial_rendtrib_total', 'dif_rendtrib',
        'sistema_prev_total', 'esocial_prev_total', 'dif_previdencia',
        'sistema_ded_dependentes', 'esocial_ded_dependentes', 'dif_dependentes',
        'sistema_plano_saude', 'esocial_plano_saude', 'dif_plano_saude',
        'sistema_isentos_total', 'esocial_isentos_total', 'dif_isentos',
        'sistema_irrf_total', 'esocial_irrf_total', 'dif_irrf'
    ]
    df_all_exp = df[cols_all].copy()
    df_all_exp.columns = [
        'CÓD. EMP', 'RAZÃO SOCIAL', 'CNPJ', 'COMPETÊNCIA', 'CPF', 'COLABORADOR',
        'STATUS AUDITORIA', 'DETALHES EXPLICATIVOS',
        'REND. TRIB. SISTEMA', 'REND. TRIB. eSOCIAL', 'DIF. RENDIMENTOS',
        'INSS SISTEMA', 'INSS eSOCIAL', 'DIF. INSS',
        'DEP. SISTEMA', 'DEP. eSOCIAL', 'DIF. DEPENDENTES',
        'PLANO SAÚDE SISTEMA', 'PLANO SAÚDE eSOCIAL', 'DIF. PLANO SAÚDE',
        'ISENTOS SISTEMA', 'ISENTOS eSOCIAL', 'DIF. ISENTOS',
        'IRRF SISTEMA', 'IRRF eSOCIAL', 'DIF. IRRF'
    ]

    ws_all = wb.create_sheet(title="📋 Conciliação Completa")
    ws_all.append(list(df_all_exp.columns))
    for r in df_all_exp.itertuples(index=False):
        ws_all.append(list(r))
    aplicar_estilos_gerais(ws_all, "Conciliacao Completa")

    wb.save(output_path)
    print(f"Planilha Executiva Auditada gerada com sucesso: {output_path}")
    print("=" * 85)


def main():
    import time
    ano_audit = 2026

    t0 = time.time()
    df_resultado = extrair_e_auditar_dados(ano=ano_audit)
    print(f"[TEMPO] Extração + cruzamento (banco): {time.time() - t0:.1f}s")

    if df_resultado is not None:
        t1 = time.time()
        gerar_relatorio_executivo_excel(df_resultado, "relatorios/Auditoria_Extrator_DIRF_eSocial_2026.xlsx")
        print(f"[TEMPO] Geração Excel #1: {time.time() - t1:.1f}s")

        t2 = time.time()
        gerar_relatorio_executivo_excel(df_resultado, "relatorios/Comparacao_Valor_Sistema_eSocial.xlsx")
        print(f"[TEMPO] Geração Excel #2: {time.time() - t2:.1f}s")

        print(f"[TEMPO] Total: {time.time() - t0:.1f}s")


if __name__ == "__main__":
    main()
