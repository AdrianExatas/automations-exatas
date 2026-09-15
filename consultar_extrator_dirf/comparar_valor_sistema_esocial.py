"""
=============================================================================
CONCILIAÇÃO OFICIAL: VALOR SISTEMA x VALOR eSOCIAL (EXTRATOR DA DIRF)
=============================================================================
Reconcilia os valores calculados pelo Domínio (regime de caixa) com os
valores retornados pelo eSocial (S-5002) processados pelo Extrator da DIRF.
Gera relatório sintético e analítico com apuração exata das diferenças.
=============================================================================
"""

import argparse
import os
import sys
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_env, get_connection


def aplicar_estilos_excel(ws, title, is_summary=False):
    # Cores modernas
    cor_header_bg = "1F4E79"       # Azul escuro
    cor_header_fg = "FFFFFF"       # Branco
    cor_zebra = "F2F5F9"           # Cinza azulado muito claro
    cor_alerta_bg = "FCE4D6"       # Salmão / Laranja suave
    cor_alerta_fg = "C00000"       # Vermelho escuro
    cor_ok_bg = "E2EFDA"           # Verde suave
    cor_ok_fg = "375623"           # Verde escuro

    font_header = Font(name="Calibri", size=11, bold=True, color=cor_header_fg)
    fill_header = PatternFill(start_color=cor_header_bg, end_color=cor_header_bg, fill_type="solid")
    border_thin = Side(border_style="thin", color="D9D9D9")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    # Congelar painel de cabeçalho
    ws.freeze_panes = "A2"

    # Estilizar cabeçalho
    for cell in ws[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = box_border

    ws.row_dimensions[1].height = 28

    # Formatar linhas de dados
    num_cols = ws.max_column
    num_rows = ws.max_row

    for row_idx in range(2, num_rows + 1):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)

        # Checar status da linha se for planilha de resumo ou detalhe
        status_val = ""
        for c in range(1, num_cols + 1):
            val = str(ws.cell(row=row_idx, column=c).value or "").strip().upper()
            if "DIVERG" in val or "PENDENTE" in val:
                status_val = "DIVERGENCIA"
            elif val == "OK" or "CONCILIADO" in val:
                if status_val != "DIVERGENCIA":
                    status_val = "OK"

        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = box_border

            # Zebra padrão
            if is_even:
                cell.fill = PatternFill(start_color=cor_zebra, end_color=cor_zebra, fill_type="solid")

            # Formatação numérica
            val = cell.value
            header_name = str(ws.cell(row=1, column=col_idx).value or "").upper()

            if isinstance(val, (int, float, Decimal)):
                if any(k in header_name for k in ["VALOR", "REND", "PREV", "IRRF", "DIF", "TOTAL", "ISENT", "SAUDE", "DEP"]):
                    cell.number_format = '"R$ "#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    # Destacar diferenças diferentes de zero
                    if "DIF" in header_name and abs(float(val)) > 0.01:
                        cell.font = Font(name="Calibri", size=10, bold=True, color=cor_alerta_fg)
                        cell.fill = PatternFill(start_color=cor_alerta_bg, end_color=cor_alerta_bg, fill_type="solid")
                elif "QTD" in header_name or "COLABORADORES" in header_name or "COD" in header_name:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            elif isinstance(val, (datetime, date)):
                cell.number_format = 'yyyy-mm'
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                str_v = str(val or "")
                if str_v in ["COM DIVERGÊNCIA", "DIVERGENTE", "PENDENTE"]:
                    cell.font = Font(name="Calibri", size=10, bold=True, color=cor_alerta_fg)
                    cell.fill = PatternFill(start_color=cor_alerta_bg, end_color=cor_alerta_bg, fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif str_v in ["CONCILIADO", "OK"]:
                    cell.font = Font(name="Calibri", size=10, bold=True, color=cor_ok_fg)
                    cell.fill = PatternFill(start_color=cor_ok_bg, end_color=cor_ok_bg, fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif any(k in header_name for k in ["CPF", "CNPJ", "DATA", "COMP"]):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-ajuste de largura
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                if isinstance(val, (datetime, date)):
                    s = "2026-00-00"
                elif isinstance(val, (int, float, Decimal)):
                    s = f"R$ {val:,.2f}"
                else:
                    s = str(val)
                max_len = max(max_len, len(s))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Ativar autofiltro
    ws.auto_filter.ref = ws.dimensions


def executar_conciliacao(codi_emp=None, ano=None, competencia=None, output_path=None):
    if not output_path:
        output_path = Path("relatorios") / "Comparacao_Valor_Sistema_eSocial.xlsx"
    else:
        output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    print("=" * 80)
    print("INICIANDO CONCILIACAO: VALOR DO SISTEMA x VALOR eSOCIAL")
    print("=" * 80)
    if codi_emp:
        print(f"Filtro Empresa:     {codi_emp}")
    if ano:
        print(f"Filtro Ano:         {ano}")
    if competencia:
        print(f"Filtro Competência: {competencia}")
    print(f"Arquivo de Saída:   {output_path}")
    print("-" * 80)

    conn = get_connection()
    cursor = conn.cursor()

    # 1. Filtro base
    where_es_parts = []
    if codi_emp:
        where_es_parts.append(f"tot.CODI_EMP = {codi_emp}")
    if competencia:
        where_es_parts.append(f"tot.COMPETENCIA_REF = '{competencia}'")
    elif ano:
        where_es_parts.append(f"YEAR(tot.COMPETENCIA_REF) = {ano}")
    where_es = ("AND " + " AND ".join(where_es_parts)) if where_es_parts else ""

    # 2. Consultar eSocial (S-5002 Totalizador)
    print("\n[1/5] Consultando eSocial S-5002 Extrator Totalizador...")
    sql_es = f"""
        SELECT 
            tot.CODI_EMP,
            emp.NOME_EMP,
            tot.COMPETENCIA_REF,
            ext.CPF,
            COALESCE(emp_col.NOME, 'COLABORADOR CPF ' + ext.CPF) AS NOME_COLABORADOR,
            COALESCE(tot.VLRRENDTRIB, 0)        AS ESOCIAL_RENDTRIB_MENSAL,
            COALESCE(tot.VLRRENDTRIB13, 0)      AS ESOCIAL_RENDTRIB_13,
            COALESCE(tot.VLRPREVOFICIAL, 0)     AS ESOCIAL_PREV_MENSAL,
            COALESCE(tot.VLRPREVOFICIAL13, 0)   AS ESOCIAL_PREV_13,
            COALESCE(tot.VLRCRMEN, 0)           AS ESOCIAL_IRRF_MENSAL,
            COALESCE(tot.VLRCR13MEN, 0)         AS ESOCIAL_IRRF_13,
            COALESCE(tot.VLRPARCISENTA65, 0) + COALESCE(tot.VLRPARCISENTA65DEC, 0) +
            COALESCE(tot.VLRDIARIAS, 0) + COALESCE(tot.VLRAJUDACUSTO, 0) +
            COALESCE(tot.VLRINDRESCONTRATO, 0) + COALESCE(tot.VLRABONOPEC, 0) +
            COALESCE(tot.VLRRENDMOLEGRAVE, 0) + COALESCE(tot.VLRRENDMOLEGRAVE13, 0) +
            COALESCE(tot.VLRAUXMORADIA, 0) + COALESCE(tot.VLRISENOUTROS, 0) AS ESOCIAL_ISENTOS_TOTAL,
            tot.I_DADOS_EVENTOS
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_TOTALIZADOR tot
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = tot.CODI_EMP AND ext.I_DADOS_EVENTOS = tot.I_DADOS_EVENTOS
        JOIN bethadba.GEEMPRE emp 
          ON emp.CODI_EMP = tot.CODI_EMP
        LEFT JOIN bethadba.FOEMPREGADOS emp_col
          ON emp_col.CODI_EMP = tot.CODI_EMP AND emp_col.CPF = ext.CPF
        JOIN bethadba.FOESOCIAL_DADOS_EVENTOS fde
          ON fde.CODI_EMP = tot.CODI_EMP AND fde.I_DADOS_EVENTOS = tot.I_DADOS_EVENTOS
        WHERE fde.VALIDADO = 1
          AND NOT EXISTS (
              SELECT 1 FROM bethadba.FOESOCIAL_DADOS_EVENTOS ret
              WHERE ret.CODI_EMP_EVENTO_RETIFICADO = fde.CODI_EMP
                AND ret.I_DADOS_EVENTOS_EVENTO_RETIFICADO = fde.I_DADOS_EVENTOS
                AND ret.VALIDADO = 1
          )
          {where_es}
    """
    cursor.execute(sql_es)
    df_es = pd.DataFrame.from_records(cursor.fetchall(), columns=[c[0].lower() for c in cursor.description])
    print(f"      Total registros eSocial retornados: {len(df_es)}")

    if df_es.empty:
        print("Nenhum dado encontrado para os filtros informados.")
        conn.close()
        return

    # 3. Subtabelas do eSocial (Dependentes e Plano de Saúde)
    print("\n[2/5] Consultando subtabelas eSocial (Dependentes e Plano de Saúde)...")
    where_dep = where_es.replace("tot.", "ded.")
    sql_dep = f"""
        SELECT ded.CODI_EMP, ded.COMPETENCIA_REF, ext.CPF,
               SUM(COALESCE(ded.VLRDEDDEP, 0)) AS ESOCIAL_DED_DEPENDENTES
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_DED_DEPEN ded
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = ded.CODI_EMP AND ext.I_DADOS_EVENTOS = ded.I_DADOS_EVENTOS
        WHERE 1=1 {where_dep}
        GROUP BY ded.CODI_EMP, ded.COMPETENCIA_REF, ext.CPF
    """
    cursor.execute(sql_dep)
    df_dep = pd.DataFrame.from_records(cursor.fetchall(), columns=[c[0].lower() for c in cursor.description])

    where_ps = where_es.replace("tot.", "ps.")
    sql_ps = f"""
        SELECT ps.CODI_EMP, ps.COMPETENCIA_REF, ext.CPF,
               SUM(COALESCE(ps.VLRSAUDETIT, 0)) AS ESOCIAL_PLANO_SAUDE
        FROM bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR_PLANO_SAUDE ps
        JOIN bethadba.FOESOCIAL_ARQUIVO_RETORNO_S_5002_EXTRATOR ext
          ON ext.CODI_EMP = ps.CODI_EMP AND ext.I_DADOS_EVENTOS = ps.I_DADOS_EVENTOS
        WHERE 1=1 {where_ps}
        GROUP BY ps.CODI_EMP, ps.COMPETENCIA_REF, ext.CPF
    """
    cursor.execute(sql_ps)
    df_ps = pd.DataFrame.from_records(cursor.fetchall(), columns=[c[0].lower() for c in cursor.description])

    # Consolidar eSocial por (CODI_EMP, COMPETENCIA_REF, CPF)
    df_es_grp = df_es.groupby(['codi_emp', 'competencia_ref', 'cpf']).agg({
        'nome_emp': 'first',
        'nome_colaborador': 'first',
        'esocial_rendtrib_mensal': 'sum',
        'esocial_rendtrib_13': 'sum',
        'esocial_prev_mensal': 'sum',
        'esocial_prev_13': 'sum',
        'esocial_irrf_mensal': 'sum',
        'esocial_irrf_13': 'sum',
        'esocial_isentos_total': 'sum'
    }).reset_index()

    if not df_dep.empty:
        df_es_grp = pd.merge(df_es_grp, df_dep, on=['codi_emp', 'competencia_ref', 'cpf'], how='left')
    else:
        df_es_grp['esocial_ded_dependentes'] = 0.0

    if not df_ps.empty:
        df_es_grp = pd.merge(df_es_grp, df_ps, on=['codi_emp', 'competencia_ref', 'cpf'], how='left')
    else:
        df_es_grp['esocial_plano_saude'] = 0.0

    df_es_grp = df_es_grp.fillna(0.0)

    # 4. Consultar Sistema (Bases de Cálculo do IRRF e Folha)
    print("\n[3/5] Consultando Bases de Cálculo do Sistema (FOBASESSERVIRRF + FOBASESIRRF)...")
    empresas = list(df_es['codi_emp'].unique())
    emp_str = ",".join(map(str, empresas))
    
    # Lista de competências em escopo
    comp_dates = sorted(list(df_es['competencia_ref'].unique()))
    comp_min = comp_dates[0].strftime("%Y-%m-%d")
    comp_max = comp_dates[-1].strftime("%Y-%m-%d")

    sql_sis = f"""
        SELECT 
            s.CODI_EMP,
            YMD(YEAR(b.DATA_PAGTO), MONTH(b.DATA_PAGTO), 1) AS COMPETENCIA_REF,
            e.CPF,
            e.NOME                                          AS NOME_COLABORADOR,
            SUM(CASE WHEN b.TIPO NOT IN (4, 6) THEN b.BASE ELSE 0 END) AS SISTEMA_RENDTRIB_MENSAL,
            SUM(CASE WHEN b.TIPO = 6 THEN b.BASE ELSE 0 END)           AS SISTEMA_RENDTRIB_13,
            SUM(CASE WHEN b.TIPO = 4 THEN b.BASE ELSE 0 END)           AS SISTEMA_RENDTRIB_PLR,
            SUM(CASE WHEN b.TIPO NOT IN (4, 6) THEN b.ABATIMENTOS ELSE 0 END) AS SISTEMA_PREV_MENSAL,
            SUM(CASE WHEN b.TIPO = 6 THEN b.ABATIMENTOS ELSE 0 END)           AS SISTEMA_PREV_13,
            SUM(b.DEPEND_DESCONTO)                                            AS SISTEMA_DED_DEPENDENTES,
            SUM(CASE WHEN b.TIPO NOT IN (4, 6) THEN b.VALOR ELSE 0 END) AS SISTEMA_IMPOSTO_RETIDO_MENSAL,
            SUM(CASE WHEN b.TIPO = 6 THEN b.VALOR ELSE 0 END)           AS SISTEMA_IMPOSTO_RETIDO_13,
            SUM(CASE WHEN b.TIPO = 4 THEN b.VALOR ELSE 0 END)           AS SISTEMA_IMPOSTO_RETIDO_PLR
        FROM bethadba.FOBASESSERVIRRF s
        JOIN bethadba.FOBASESIRRF b 
          ON b.I_BASESIRRF = s.I_BASESIRRF
        JOIN bethadba.FOEMPREGADOS e 
          ON e.CODI_EMP = s.CODI_EMP 
         AND e.I_EMPREGADOS = s.I_EMPREGADOS
        WHERE s.CODI_EMP IN ({emp_str})
          AND s.RATEIO = 0
          AND b.DATA_PAGTO >= '{comp_min}' 
          AND b.DATA_PAGTO <= DATEADD(month, 1, '{comp_max}') - 1
        GROUP BY 
            s.CODI_EMP,
            YMD(YEAR(b.DATA_PAGTO), MONTH(b.DATA_PAGTO), 1),
            e.CPF,
            e.NOME
    """
    cursor.execute(sql_sis)
    df_sis = pd.DataFrame.from_records(cursor.fetchall(), columns=[c[0].lower() for c in cursor.description])
    print(f"      Total registros Sistema retornados: {len(df_sis)}")

    # 5. Consultar Eventos Complementares do Sistema (Plano de Saúde e Isentos)
    print("\n[4/5] Consultando Eventos de Plano de Saúde e Isentos do Sistema...")
    sql_eve = f"""
        SELECT 
            bs.CODI_EMP,
            YMD(YEAR(bs.DATA_PAGTO), MONTH(bs.DATA_PAGTO), 1) AS COMPETENCIA_REF,
            e.CPF,
            SUM(CASE WHEN ev.CLASSIFICACAO = 29 THEN
                    CASE WHEN m.PROV_DESC IN ('D', 'ID') THEN m.VALOR_CAL ELSE -m.VALOR_CAL END
                ELSE 0 END) AS SISTEMA_PLANO_SAUDE,
            SUM(CASE WHEN ev.REND_ISENTOS > 0 OR ev.CLASSIFICACAO = 46 
                     OR ev.CODIGO_INCIDENCIA_IRRF_ESOCIAL IN (72, 73) THEN
                    CASE WHEN m.PROV_DESC IN ('P', 'I') THEN m.VALOR_CAL ELSE -m.VALOR_CAL END
                ELSE 0 END) AS SISTEMA_ISENTOS_TOTAL
        FROM bethadba.FOBASESSERV bs
        JOIN bethadba.FOMOVTOSERV m 
          ON m.CODI_EMP = bs.CODI_EMP AND m.I_CALCULOS = bs.I_CALCULOS
        JOIN bethadba.FOEMPREGADOS e 
          ON e.CODI_EMP = bs.CODI_EMP AND e.I_EMPREGADOS = bs.I_EMPREGADOS
        JOIN bethadba.FOPARMTO p 
          ON p.CODI_EMP = bs.CODI_EMP
        JOIN bethadba.FOEVENTOS ev 
          ON ev.CODI_EMP = p.CODI_EMP_EVE AND ev.I_EVENTOS = m.I_EVENTOS
        WHERE bs.CODI_EMP IN ({emp_str})
          AND m.RATEIO = 0
          AND m.ORIGEM <> 'F'
          AND ev.SOMA_INF_REN = 'S'
          AND bs.DATA_PAGTO >= '{comp_min}' 
          AND bs.DATA_PAGTO <= DATEADD(month, 1, '{comp_max}') - 1
        GROUP BY 
            bs.CODI_EMP,
            YMD(YEAR(bs.DATA_PAGTO), MONTH(bs.DATA_PAGTO), 1),
            e.CPF
    """
    cursor.execute(sql_eve)
    df_eve = pd.DataFrame.from_records(cursor.fetchall(), columns=[c[0].lower() for c in cursor.description])
    conn.close()

    if not df_eve.empty:
        df_sis = pd.merge(df_sis, df_eve, on=['codi_emp', 'competencia_ref', 'cpf'], how='left')
    else:
        df_sis['sistema_plano_saude'] = 0.0
        df_sis['sistema_isentos_total'] = 0.0

    df_sis = df_sis.fillna(0.0)

    # 6. Realizar o Cruzamento e Conciliação
    print("\n[5/5] Realizando Cruzamento e Análise de Divergências...")
    conciliado = pd.merge(df_es_grp, df_sis, on=['codi_emp', 'competencia_ref', 'cpf'], how='outer')

    # Ajustar nomes de colaboradores e empresas nulos
    conciliado['nome_emp'] = conciliado['nome_emp'].fillna("EMPRESA " + conciliado['codi_emp'].astype(str))
    conciliado['nome_colaborador'] = conciliado['nome_colaborador_x'].combine_first(conciliado['nome_colaborador_y']).fillna("N/D")
    conciliado.drop(columns=['nome_colaborador_x', 'nome_colaborador_y'], errors='ignore', inplace=True)

    # Preencher numéricos vazios e converter para float
    cols_num = [c for c in conciliado.columns if c.startswith(('esocial_', 'sistema_'))]
    for c in cols_num:
        conciliado[c] = pd.to_numeric(conciliado[c], errors='coerce').fillna(0.0).astype(float)

    # Totais consolidados
    conciliado['esocial_rendtrib_total'] = (conciliado['esocial_rendtrib_mensal'] + conciliado['esocial_rendtrib_13']).round(2)
    conciliado['sistema_rendtrib_total'] = (conciliado['sistema_rendtrib_mensal'] + conciliado['sistema_rendtrib_13'] + conciliado['sistema_rendtrib_plr']).round(2)

    conciliado['esocial_prev_total'] = (conciliado['esocial_prev_mensal'] + conciliado['esocial_prev_13']).round(2)
    conciliado['sistema_prev_total'] = (conciliado['sistema_prev_mensal'] + conciliado['sistema_prev_13']).round(2)

    conciliado['esocial_irrf_total'] = (conciliado['esocial_irrf_mensal'] + conciliado['esocial_irrf_13']).round(2)
    conciliado['sistema_irrf_total'] = (conciliado['sistema_imposto_retido_mensal'] + conciliado['sistema_imposto_retido_13'] + conciliado['sistema_imposto_retido_plr']).round(2)

    # Diferenças individuais
    conciliado['dif_rendtrib'] = (conciliado['sistema_rendtrib_total'] - conciliado['esocial_rendtrib_total']).round(2)
    conciliado['dif_previdencia'] = (conciliado['sistema_prev_total'] - conciliado['esocial_prev_total']).round(2)
    conciliado['dif_dependentes'] = (conciliado['sistema_ded_dependentes'] - conciliado['esocial_ded_dependentes']).round(2)
    conciliado['dif_plano_saude'] = (conciliado['sistema_plano_saude'] - conciliado['esocial_plano_saude']).round(2)
    conciliado['dif_irrf'] = (conciliado['sistema_irrf_total'] - conciliado['esocial_irrf_total']).round(2)
    conciliado['dif_isentos'] = (conciliado['sistema_isentos_total'] - conciliado['esocial_isentos_total']).round(2)

    # Diferença absoluta total
    conciliado['dif_absoluta_total'] = (
        conciliado['dif_rendtrib'].abs() +
        conciliado['dif_previdencia'].abs() +
        conciliado['dif_dependentes'].abs() +
        conciliado['dif_plano_saude'].abs() +
        conciliado['dif_irrf'].abs() +
        conciliado['dif_isentos'].abs()
    ).round(2)

    # Status de conciliação
    def definir_status(row):
        if row['dif_absoluta_total'] > 0.01:
            return "COM DIVERGÊNCIA"
        return "CONCILIADO"

    conciliado['status'] = conciliado.apply(definir_status, axis=1)

    # Descrever motivos
    def descrever_motivo(r):
        if r['status'] == "CONCILIADO":
            return "Conciliado perfeitamente (Sem divergência)"
        motivos = []
        if abs(r['dif_rendtrib']) > 0.01:
            motivos.append(f"Rend. Trib: Sis R$ {r['sistema_rendtrib_total']:,.2f} x eSoc R$ {r['esocial_rendtrib_total']:,.2f} (Dif R$ {r['dif_rendtrib']:,.2f})")
        if abs(r['dif_previdencia']) > 0.01:
            motivos.append(f"Prev. Oficial: Sis R$ {r['sistema_prev_total']:,.2f} x eSoc R$ {r['esocial_prev_total']:,.2f} (Dif R$ {r['dif_previdencia']:,.2f})")
        if abs(r['dif_dependentes']) > 0.01:
            motivos.append(f"Dependentes: Sis R$ {r['sistema_ded_dependentes']:,.2f} x eSoc R$ {r['esocial_ded_dependentes']:,.2f} (Dif R$ {r['dif_dependentes']:,.2f})")
        if abs(r['dif_plano_saude']) > 0.01:
            motivos.append(f"Plano Saúde: Sis R$ {r['sistema_plano_saude']:,.2f} x eSoc R$ {r['esocial_plano_saude']:,.2f} (Dif R$ {r['dif_plano_saude']:,.2f})")
        if abs(r['dif_irrf']) > 0.01:
            motivos.append(f"IRRF Retido: Sis R$ {r['sistema_irrf_total']:,.2f} x eSoc R$ {r['esocial_irrf_total']:,.2f} (Dif R$ {r['dif_irrf']:,.2f})")
        if abs(r['dif_isentos']) > 0.01:
            motivos.append(f"Isentos: Sis R$ {r['sistema_isentos_total']:,.2f} x eSoc R$ {r['esocial_isentos_total']:,.2f} (Dif R$ {r['dif_isentos']:,.2f})")
        return " | ".join(motivos)

    conciliado['detalhes_divergencia'] = conciliado.apply(descrever_motivo, axis=1)

    # Estatísticas
    total_linhas = len(conciliado)
    total_conciliados = (conciliado['status'] == "CONCILIADO").sum()
    total_divergentes = (conciliado['status'] == "COM DIVERGÊNCIA").sum()

    print(f"\nResultado da Conciliação:")
    print(f"  Total Colaboradores / Competências: {total_linhas}")
    print(f"  Conciliados (100% batimento):       {total_conciliados} ({total_conciliados/total_linhas*100:.1f}%)")
    print(f"  Com divergência:                    {total_divergentes} ({total_divergentes/total_linhas*100:.1f}%)")

    # 7. Resumo por Empresa
    resumo_emp = conciliado.groupby(['codi_emp', 'nome_emp']).agg({
        'cpf': 'nunique',
        'status': lambda s: (s == "COM DIVERGÊNCIA").sum(),
        'sistema_rendtrib_total': 'sum',
        'esocial_rendtrib_total': 'sum',
        'dif_rendtrib': 'sum',
        'sistema_prev_total': 'sum',
        'esocial_prev_total': 'sum',
        'dif_previdencia': 'sum',
        'sistema_irrf_total': 'sum',
        'esocial_irrf_total': 'sum',
        'dif_irrf': 'sum'
    }).reset_index()

    resumo_emp.rename(columns={
        'codi_emp': 'CODI_EMP',
        'nome_emp': 'NOME_EMP',
        'cpf': 'QTD_COLABORADORES',
        'status': 'QTD_DIVERGENCIAS',
        'sistema_rendtrib_total': 'SISTEMA_RENDTRIB',
        'esocial_rendtrib_total': 'ESOCIAL_RENDTRIB',
        'dif_rendtrib': 'DIF_RENDTRIB',
        'sistema_prev_total': 'SISTEMA_PREV',
        'esocial_prev_total': 'ESOCIAL_PREV',
        'dif_previdencia': 'DIF_PREV',
        'sistema_irrf_total': 'SISTEMA_IRRF',
        'esocial_irrf_total': 'ESOCIAL_IRRF',
        'dif_irrf': 'DIF_IRRF'
    }, inplace=True)

    resumo_emp['SITUACAO'] = resumo_emp['QTD_DIVERGENCIAS'].apply(
        lambda q: "COM DIVERGÊNCIA" if q > 0 else "OK"
    )
    resumo_emp.sort_values(by=['QTD_DIVERGENCIAS', 'CODI_EMP'], ascending=[False, True], inplace=True)

    # 8. Aba de Divergências Detalhadas
    df_div = conciliado[conciliado['status'] == "COM DIVERGÊNCIA"].copy()
    colunas_detalhe = [
        'codi_emp', 'nome_emp', 'competencia_ref', 'cpf', 'nome_colaborador',
        'status', 'detalhes_divergencia',
        'sistema_rendtrib_total', 'esocial_rendtrib_total', 'dif_rendtrib',
        'sistema_prev_total', 'esocial_prev_total', 'dif_previdencia',
        'sistema_ded_dependentes', 'esocial_ded_dependentes', 'dif_dependentes',
        'sistema_plano_saude', 'esocial_plano_saude', 'dif_plano_saude',
        'sistema_irrf_total', 'esocial_irrf_total', 'dif_irrf',
        'sistema_isentos_total', 'esocial_isentos_total', 'dif_isentos'
    ]

    df_div_export = df_div[colunas_detalhe].copy()
    df_div_export.columns = [c.upper() for c in df_div_export.columns]

    # Aba de Conciliação Completa
    df_all_export = conciliado[colunas_detalhe].copy()
    df_all_export.columns = [c.upper() for c in df_all_export.columns]

    # 9. Escrever na Planilha Excel
    print(f"\nGerando planilha formatada em {output_path}...")
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Aba 1: Resumo Empresas
    ws1 = wb.create_sheet(title="Resumo Empresas")
    ws1.append(list(resumo_emp.columns))
    for r in resumo_emp.itertuples(index=False):
        ws1.append(list(r))
    aplicar_estilos_excel(ws1, "Resumo Empresas", is_summary=True)

    # Aba 2: Divergências
    ws2 = wb.create_sheet(title="Divergencias")
    ws2.append(list(df_div_export.columns))
    for r in df_div_export.itertuples(index=False):
        ws2.append(list(r))
    aplicar_estilos_excel(ws2, "Divergencias")

    # Aba 3: Conciliação Completa
    ws3 = wb.create_sheet(title="Conciliacao Completa")
    ws3.append(list(df_all_export.columns))
    for r in df_all_export.itertuples(index=False):
        ws3.append(list(r))
    aplicar_estilos_excel(ws3, "Conciliacao Completa")

    wb.save(output_path)
    print(f"Planilha gerada com sucesso: {output_path}")
    print("=" * 80)


def main():
    parser = argparse.ArgumentParser(description="Compara Valor do Sistema x Valor eSocial no Extrator da DIRF.")
    parser.add_argument("--codi-emp", type=int, help="Código da Empresa a auditar")
    parser.add_argument("--ano", type=int, help="Ano de competência (ex: 2026)")
    parser.add_argument("--competencia", type=str, help="Competência específica (ex: 2026-01-01)")
    parser.add_argument("--output", type=str, default="relatorios/Comparacao_Valor_Sistema_eSocial.xlsx", help="Caminho do arquivo Excel de saída")
    args = parser.parse_args()

    executar_conciliacao(
        codi_emp=args.codi_emp,
        ano=args.ano,
        competencia=args.competencia,
        output_path=args.output
    )


if __name__ == "__main__":
    main()
